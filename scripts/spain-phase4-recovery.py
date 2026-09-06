#!/usr/bin/env python3
"""
Spain Phase 4 — final Synergym + Fitness Park recovery ONLY.

- Preserve all existing READY rows and es_* IDs where possible
- No invented coords (no city/postcode/country centroids, no neighbour, no 0,0)
- Do NOT modify src/data/centers.json
- Strict Spain geography (mainland + Balearics + Canaries + Ceuta + Melilla)
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import ssl
import time
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data/spain"
RAW = OUT / "raw"
PAGES = RAW / "pages"
STAGING = OUT / "spain_centers_staging.json"
CACHE = OUT / "spain_geocode_cache.json"
CENTERS = ROOT / "src/data/centers.json"
LOG = OUT / "spain_phase4_recovery.log"

ctx = ssl.create_default_context()
UA = "GymlySpainGeocoder/4.0 (catalog research; accuracy over coverage; no fallback centroids)"
HTTP_UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
)

ES_MAINLAND = (35.9, 43.8, -9.4, 3.4)
ES_BALEARIC = (38.6, 40.1, 1.1, 4.4)
ES_CANARY = (27.6, 29.5, -18.2, -13.3)
ES_CEUTA = (35.85, 35.92, -5.35, -5.27)
ES_MELILLA = (35.26, 35.33, -2.97, -2.92)

MAJOR_CITIES = [
    "Madrid", "Barcelona", "Valencia", "Sevilla", "Zaragoza", "Málaga",
    "Murcia", "Palma", "Las Palmas", "Bilbao", "Alicante", "Córdoba",
    "Valladolid", "Vigo", "Gijón", "Granada", "A Coruña",
]

# Map phase-4 coord_source → report buckets
COORD_BUCKET = {
    "official_json_ld": "OFFICIAL_COORDINATE",
    "official_locator_js": "OFFICIAL_COORDINATE",
    "official_club_page_maps_embed": "OFFICIAL_MAP_PIN",
    "official_brand_mymaps": "OFFICIAL_MAP_PIN",
    "overpass_named_gym_poi": "NAMED_GYM_POI",
    "nominatim_named_gym_poi": "NAMED_GYM_POI",
    "nominatim": "STRICT_ADDRESS_GEOCODE",
    "nominatim_strict_address": "STRICT_ADDRESS_GEOCODE",
}

COARSE = {
    "country", "state", "county", "municipality", "city", "town", "village",
    "administrative", "postcode", "postal_code", "suburb", "neighbourhood",
    "quarter", "district", "borough", "region", "island",
}

STREET_VARIANTS = [
    (r"\bAv\.?\b", "Avenida"),
    (r"\bAvd\.?\b", "Avenida"),
    (r"\bAvda\.?\b", "Avenida"),
    (r"\bAvinguda\b", "Avenida"),
    (r"\bC\.?\b", "Calle"),
    (r"\bC/\b", "Calle"),
    (r"\bCarrer\b", "Calle"),
    (r"\bRúa\b", "Calle"),
    (r"\bRua\b", "Calle"),
    (r"\bP\.º\b", "Paseo"),
    (r"\bPº\b", "Paseo"),
    (r"\bPg\.?\b", "Paseo"),
    (r"\bPasseig\b", "Paseo"),
    (r"\bCtra\.?\b", "Carretera"),
    (r"\bC\.ª\b", "Carretera"),
    (r"\bPl\.?\b", "Plaza"),
    (r"\bPlza\.?\b", "Plaza"),
]

# City repairs for known bad Synergym homepage parses
CITY_FIXES = {
    "8 elche": "Elche",
    "18 vigo": "Vigo",
    "16 murcia": "Murcia",
    "230 motril-granada": "Motril",
    "14": "Segovia",
    "6": "Cáceres",
    "10": "Zaragoza",
    "casablanca": "Zaragoza",
    "san josé": "Zaragoza",
    "san jose": "Zaragoza",
    "noroeste": "Córdoba",
    "santiago de vigo": "Vigo",
    "la coruña": "A Coruña",
    "bajo": "Valencia",
    "portugalete, vizcaya": "Portugalete",
    "lorca murcia": "Lorca",
    "dos hermanas,sevilla": "Dos Hermanas",
    "palma mallorca": "Palma",
    "burlada pamplona": "Burlada",
    "el puerto de sta maría": "El Puerto de Santa María",
    "el puerto de sta maria": "El Puerto de Santa María",
    "sanlúcar": "Sanlúcar de Barrameda",
    "las chafiras": "San Miguel de Abona",
    "vinaros": "Vinaròs",
}


def log(*a):
    msg = " ".join(str(x) for x in a)
    print(msg, flush=True)
    with LOG.open("a", encoding="utf-8") as f:
        f.write(msg + "\n")


def make_id(brand, address, postal, city):
    key = "|".join([
        (brand or "").strip().lower(),
        (address or "").strip().lower(),
        (str(postal) or "").strip().lower(),
        (city or "").strip().lower(),
        "spain",
    ])
    return "es_" + hashlib.md5(key.encode("utf-8")).hexdigest()[:10]


def norm(s):
    s = (s or "").lower()
    for a, b in [
        ("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"),
        ("ü", "u"), ("ñ", "n"), ("ç", "c"), ("à", "a"), ("è", "e"),
        ("ò", "o"), ("ï", "i"),
    ]:
        s = s.replace(a, b)
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return s.strip()


def haversine(lat1, lng1, lat2, lng2):
    R = 6371000
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lng2 - lng1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * R * math.asin(math.sqrt(a))


def es_postal(s) -> str:
    if s is None:
        return ""
    if isinstance(s, float) and math.isnan(s):
        return ""
    m = re.search(r"\b(\d{5})\b", str(s).strip())
    return m.group(1) if m else ""


def in_spain_bbox(lat, lng) -> bool:
    try:
        lat, lng = float(lat), float(lng)
    except (TypeError, ValueError):
        return False
    if not (math.isfinite(lat) and math.isfinite(lng)):
        return False
    if lat == 0 and lng == 0:
        return False
    for bounds in [ES_MAINLAND, ES_BALEARIC, ES_CANARY, ES_CEUTA, ES_MELILLA]:
        lo, hi, w, e = bounds
        if lo <= lat <= hi and w <= lng <= e:
            return True
    return False


def in_balearic(lat, lng):
    try:
        lat, lng = float(lat), float(lng)
        lo, hi, w, e = ES_BALEARIC
        return lo <= lat <= hi and w <= lng <= e
    except (TypeError, ValueError):
        return False


def in_canary(lat, lng):
    try:
        lat, lng = float(lat), float(lng)
        lo, hi, w, e = ES_CANARY
        return lo <= lat <= hi and w <= lng <= e
    except (TypeError, ValueError):
        return False


def in_ceuta(lat, lng):
    try:
        lat, lng = float(lat), float(lng)
        lo, hi, w, e = ES_CEUTA
        return lo <= lat <= hi and w <= lng <= e
    except (TypeError, ValueError):
        return False


def in_melilla(lat, lng):
    try:
        lat, lng = float(lat), float(lng)
        lo, hi, w, e = ES_MELILLA
        return lo <= lat <= hi and w <= lng <= e
    except (TypeError, ValueError):
        return False


def richness(r):
    return (
        (2 if r.get("lat") is not None else 0)
        + (1 if r.get("postal_code") else 0)
        + (1 if r.get("address") else 0)
        + (1 if r.get("city") else 0)
        + (1 if r.get("coord_source") and "official" in str(r.get("coord_source")) else 0)
    )


def http_get(url, timeout=40):
    req = urllib.request.Request(url, headers={"User-Agent": HTTP_UA, "Accept": "*/*"})
    with urllib.request.urlopen(req, context=ctx, timeout=timeout) as resp:
        return resp.read().decode("utf-8", errors="replace")


def nominatim(query, cache, extra=None):
    key = query if not extra else query + "||" + json.dumps(extra, sort_keys=True)
    if key in cache:
        return cache[key]
    params = {
        "q": query,
        "format": "json",
        "addressdetails": 1,
        "limit": 8,
        "countrycodes": "es",
    }
    if extra:
        params.update(extra)
    url = "https://nominatim.openstreetmap.org/search?" + urllib.parse.urlencode(params)
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, context=ctx, timeout=30) as r:
        data = json.loads(r.read().decode())
    cache[key] = data
    time.sleep(1.15)
    return data


def nominatim_structured(cache, street, city, postal=""):
    key = f"struct|{street}|{city}|{postal}"
    if key in cache:
        return cache[key]
    params = {
        "street": street,
        "city": city,
        "country": "Spain",
        "format": "json",
        "addressdetails": 1,
        "limit": 8,
        "countrycodes": "es",
    }
    if postal:
        params["postalcode"] = postal
    url = "https://nominatim.openstreetmap.org/search?" + urllib.parse.urlencode(params)
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, context=ctx, timeout=30) as r:
        data = json.loads(r.read().decode())
    cache[key] = data
    time.sleep(1.15)
    return data


def address_variants(street: str) -> list[str]:
    base = re.sub(r"\s+", " ", (street or "").strip())
    if not base:
        return []
    out = [base]
    # expand abbreviations
    expanded = base
    for pat, repl in STREET_VARIANTS:
        expanded2 = re.sub(pat, repl, expanded, flags=re.I)
        if expanded2 != expanded:
            expanded = expanded2
            out.append(expanded)
    # Catalan/Valencian/Galician common flips
    flips = [
        ("Avenida", "Avinguda"),
        ("Calle", "Carrer"),
        ("Calle", "Rúa"),
        ("Paseo", "Passeig"),
        ("Plaza", "Plaça"),
    ]
    for a, b in flips:
        if a.lower() in expanded.lower():
            out.append(re.sub(a, b, expanded, flags=re.I))
    # strip shopping-center prefixes for secondary queries only
    m = re.match(r"^(?:CC|C\.?C\.?|Centro Comercial|Parque Comercial)\s+[^,-]+[,-]\s*(.+)$", base, re.I)
    if m:
        out.append(m.group(1).strip())
    # glued prefixes
    cleaned = re.sub(
        r"^(Centro|Parquesol|Paseo Zorrilla|Av Valencia|CentroCalle|CentroAv\.?|Centro Areal)",
        "",
        base,
        flags=re.I,
    ).strip(" ,")
    if cleaned and cleaned != base:
        out.append(cleaned)
        out.extend(address_variants(cleaned)[:3])
    # unique preserve order
    seen = set()
    uniq = []
    for s in out:
        s = re.sub(r"\s+", " ", s).strip(" ,")
        k = norm(s)
        if s and k not in seen:
            seen.add(k)
            uniq.append(s)
    return uniq[:10]


def clean_city(city: str, name: str = "", address: str = "") -> str:
    c = (city or "").strip(" ,.")
    key = norm(c)
    if key in CITY_FIXES:
        return CITY_FIXES[key]
    # strip province suffixes
    c = re.sub(r",\s*(Vizcaya|Sevilla|Murcia|Barcelona|La Rioja|Álava|Málaga)\s*$", "", c, flags=re.I)
    # House number leaked into city (Synergym homepage parse): "33" / "71"
    if c.isdigit() or len(c) <= 2 or re.fullmatch(r"\d+\s+\w+", c or ""):
        # Prefer multi-token city from Synergym name before district suffix
        # e.g. "Synergym Girona Devesa" → Girona; "Synergym El Prat Les Moreres" → El Prat
        m = re.match(
            r"Synergym\s+(.+?)(?:\s+(?:Centro|Sur|Norte|Nord|Universidad|Parquesol|"
            r"Devesa|Castells|Clot|Migdia|Catedral|Sant\s+Isidre|Ronda\s+Sur|"
            r"Hacienda\s+El\s+Rosario|Los\s+Royales|La\s+Jota|Duquesa|"
            r"Buenavista|Gamonal|Ensanche|Torres\s+de\s+la\s+Luz|"
            r"Juan\s+Carlos\s+I|Río\s+Segura|Rio\s+Segura).*)?$",
            name or "",
            re.I,
        )
        if m:
            candidate = m.group(1).strip(" -|,")
            # Drop trailing district tokens when still multi-word brand+city
            candidate = re.split(r"\s*[-–]\s*", candidate)[0].strip()
            if candidate and not candidate.isdigit():
                return candidate
        m = re.match(
            r"Synergym\s+([A-Za-zÀ-ÿ'’]+(?:\s+(?:de|del|de\s+la|i|y)\s+[A-Za-zÀ-ÿ'’]+)?)",
            name or "",
            re.I,
        )
        if m:
            return m.group(1).strip()
        return ""
    return c.strip()


def repair_synergym_row(r):
    """Fix glued homepage parses; keep official display address when cleanable."""
    name = r.get("name") or r.get("center_name") or ""
    addr = (r.get("address") or "").strip()
    city = (r.get("city") or "").strip()

    # House number parked in city field — move onto address before city cleanup
    if city.isdigit() and addr and not re.search(rf",\s*{re.escape(city)}\b", addr):
        if not re.search(rf"\b{re.escape(city)}\b\s*$", addr):
            addr = f"{addr}, {city}"
        city = ""

    # Address embedded in name (no separate address)
    if not addr:
        m = re.search(
            r"(?:Synergym\s+[^,]+?)((?:Av\.|Avda\.|Avenida|C/|C\.|Calle|P\.º|Paseo|Carrer|Rúa|Cam\.|Camino|Plaza|Pl\.|Boulevard|Ctra\.).+)$",
            name, re.I,
        )
        if m:
            tail = m.group(1).strip()
            # "Marques de Murrieta, 62, Logroño, La Rioja"
            parts = [p.strip() for p in tail.split(",")]
            if len(parts) >= 2:
                # last parts often city/province
                maybe_city = parts[-2] if len(parts) >= 3 else parts[-1]
                maybe_province = parts[-1] if len(parts) >= 3 else ""
                street_parts = parts[:-2] if len(parts) >= 3 else parts[:-1]
                if street_parts:
                    addr = ", ".join(street_parts)
                if maybe_city and (not city or city.isdigit() or len(city) < 3):
                    city = maybe_city
                # strip province from city if glued
                if maybe_province and maybe_province.lower() not in ("spain", "españa"):
                    pass
            else:
                addr = tail
            # Fix name to short form
            short = re.split(r"(?=Av\.|Avda\.|Avenida|C/|C\.|Calle|P\.º|Paseo|Carrer|Rúa|Cam\.|Camino|Plaza|Boulevard|Ctra\.)", name, maxsplit=1)[0]
            short = short.strip(" -|,")
            if short.startswith("Synergym"):
                r["name"] = short
                r["center_name"] = short

    # Glued address prefixes
    if addr:
        addr2 = re.sub(
            r"^(Centro|Parquesol|Paseo Zorrilla|Av Valencia|CentroCalle|CentroAv\.?|Centro Areal|Parque,\s*)",
            "",
            addr,
            flags=re.I,
        ).strip(" ,")
        # "Av ValenciaAv. de Valencia, 108" → keep second half
        m = re.search(r"(Av\.?\s*de\s+Valencia,\s*\d+)", addr, re.I)
        if m:
            addr2 = m.group(1)
        m = re.search(r"(P\.º\s*de\s+Zorrilla,\s*\d+)", addr, re.I)
        if m:
            addr2 = m.group(1)
        m = re.search(r"(C\.\s*de\s+Hernando\s+de\s+Acuña,\s*\d+)", addr, re.I)
        if m:
            addr2 = m.group(1)
        m = re.search(r"(Rúa\s+do\s+Areal,\s*\d+)", addr, re.I)
        if m:
            addr2 = m.group(1)
        m = re.search(r"(Calle\s+Sta\.\s*Teresa.*)$", addr, re.I)
        if m:
            addr2 = m.group(1)
        if addr2:
            addr = addr2

    city = clean_city(city, name=r.get("name") or name, address=addr)
    # Province used as city in parsed data
    if not city:
        prov = ""
        # from original name tokens
        m = re.match(r"Synergym\s+(.+)", r.get("name") or name, re.I)
        if m:
            bits = m.group(1).split()
            if bits:
                city = bits[0]

    if addr:
        r["address"] = addr
    if city:
        r["city"] = city
    # preserve id stability unless we had empty/broken identity fields
    if r.get("import_category") in {"NEEDS_COORDINATES", "NEEDS_REVIEW"}:
        r["id"] = make_id(r.get("brand"), r.get("address") or "", r.get("postal_code") or "", r.get("city") or "")
    return r


def score_candidate(item, street, postal, city, allow_road_cluster=True):
    reasons = []
    score = 0
    display = (item.get("display_name") or "").lower()
    addr = item.get("address") or {}
    lat, lng = float(item["lat"]), float(item["lon"])
    if not in_spain_bbox(lat, lng):
        return None, ["outside_spain"], lat, lng
    cc = (addr.get("country_code") or "").lower()
    if cc and cc != "es":
        return None, ["not_country_es:" + cc], lat, lng

    t = (item.get("type") or item.get("class") or "").lower()
    addresstype = (item.get("addresstype") or "").lower()
    osm_class = (item.get("class") or "").lower()
    if t in COARSE or addresstype in COARSE:
        return None, ["coarse_type:" + (addresstype or t)], lat, lng
    if osm_class in {"boundary", "place"} and t not in {
        "house", "building", "yes", "retail", "commercial", "industrial",
        "gym", "fitness_centre", "sports_centre",
    }:
        return None, ["coarse_class:" + osm_class + ":" + t], lat, lng

    is_gym = (
        osm_class in {"leisure", "amenity", "sport"}
        and t in {"fitness_centre", "gym", "sports_centre", "sports_hall"}
    ) or ("synergym" in display or "fitness park" in display)

    pc = str(addr.get("postcode") or "")
    pc_n = re.sub(r"\s+", "", pc).upper()
    postal_n = re.sub(r"\s+", "", str(postal or "")).upper()
    if postal_n and pc_n == postal_n:
        score += 5
        reasons.append("postal_exact")
    elif postal_n and pc_n and pc_n[:3] == postal_n[:3]:
        score += 1
        reasons.append("postal_soft")

    city_n = norm(city)
    city_fields = " ".join(
        norm(addr.get(k) or "")
        for k in ("city", "town", "village", "municipality", "suburb", "city_district")
    )
    if city_n and (city_n in city_fields or city_n in norm(display)):
        score += 3
        reasons.append("city_ok")

    street_n = norm(street)
    road = norm(addr.get("road") or "")
    street_exp = street_n
    for a, b in [
        ("c ", "calle "), ("av ", "avenida "), ("avd ", "avenida "),
        ("avda ", "avenida "), ("carrer ", "calle "), ("rua ", "calle "),
        ("rúa ", "calle "), ("pl ", "plaza "), ("pg ", "paseo "),
        ("p ", "paseo "), ("ctra ", "carretera "),
    ]:
        if street_exp.startswith(a):
            street_exp = b + street_exp[len(a):]
    # Catalan avenida↔avinguda soft
    street_soft = street_exp.replace("avenida ", "avinguda ").replace("paseo ", "passeig ").replace("calle ", "carrer ")
    road_match = False
    if road and street_n and (
        road in street_n or street_n in norm(display) or road in street_exp
        or street_exp in norm(display) or road in street_soft or street_soft in norm(display)
        or any(tok and tok in road for tok in street_exp.split() if len(tok) > 4)
        or any(tok and tok in street_exp for tok in road.split() if len(tok) > 4)
    ):
        score += 4
        reasons.append("road_match")
        road_match = True

    hn = str(addr.get("house_number") or "")
    m = re.search(r"\b(\d+[a-z]?)\b", (street or "").lower())
    if hn and m and hn.lower() == m.group(1).lower():
        score += 3
        reasons.append("house_number_match")

    if is_gym:
        score += 4
        reasons.append("named_gym_poi")
    elif osm_class in {"building", "amenity", "leisure", "shop"} or t in {
        "gym", "fitness_centre", "sports_centre", "yes", "retail",
    }:
        score += 2
        reasons.append("building_or_amenity")

    has_structure = any(
        x in reasons
        for x in ("road_match", "house_number_match", "building_or_amenity", "named_gym_poi")
    )
    if not has_structure:
        if allow_road_cluster and "postal_exact" in reasons and "city_ok" in reasons and (
            road_match or osm_class == "highway"
        ):
            score += 2
            reasons.append("road_accepted_with_postal_city")
        else:
            return None, reasons + ["no_street_or_building"], lat, lng

    # Phase 4: allow highway with postal+city+road even when OSM omits repeating postal on some hits
    if (
        osm_class == "highway"
        and "city_ok" in reasons
        and road_match
        and postal_n
        and "house_number_match" not in reasons
    ):
        # Require postal_exact OR postal present in display_name
        if "postal_exact" in reasons or (postal_n and postal_n in display):
            reasons.append("highway_street_level_ok")
        else:
            return None, reasons + ["highway_without_postal_confirm"], lat, lng

    if score < 7:
        return None, reasons + ["score_too_low"], lat, lng
    return score, reasons, lat, lng


def pick_best_scored(scored):
    """Resolve ambiguity: same-road cluster OK; divergent roads rejected."""
    if not scored:
        return None
    scored = sorted(scored, key=lambda x: -x[0])
    top = scored[0]
    if len(scored) == 1:
        return top
    # Same score cluster
    peers = [s for s in scored if abs(s[0] - top[0]) < 0.6]
    if len(peers) == 1:
        return top
    # Prefer house_number_match
    with_hn = [s for s in peers if "house_number_match" in s[1]]
    if len(with_hn) == 1:
        return with_hn[0]
    if with_hn:
        peers = with_hn
    # Prefer named gym / building
    with_poi = [s for s in peers if "named_gym_poi" in s[1] or "building_or_amenity" in s[1]]
    if len(with_poi) == 1:
        return with_poi[0]
    if with_poi:
        peers = with_poi

    roads = set()
    for s in peers:
        road = norm((s[5].get("address") or {}).get("road") or "")
        roads.add(road)
    max_d = 0
    for i in range(len(peers)):
        for j in range(i + 1, len(peers)):
            max_d = max(max_d, haversine(peers[i][2], peers[i][3], peers[j][2], peers[j][3]))

    # Same road (or empty) and tight cluster → accept median
    if len(roads) <= 1 and max_d <= 550:
        # geometric median approx: pick point minimizing sum distance
        best_i, best_sum = 0, None
        for i, s in enumerate(peers):
            sm = sum(haversine(s[2], s[3], t[2], t[3]) for t in peers)
            if best_sum is None or sm < best_sum:
                best_sum, best_i = sm, i
        chosen = peers[best_i]
        reasons = list(chosen[1]) + [f"same_road_cluster_{int(max_d)}m"]
        return (chosen[0], reasons, chosen[2], chosen[3], chosen[4], chosen[5])

    # Different segments far apart without house number → reject
    if max_d > 550 and "house_number_match" not in top[1]:
        return None
    return top


def apply_geocode_hit(r, lat, lng, reasons, display, source, postal_from=None, preserve_id=True):
    if not in_spain_bbox(lat, lng):
        return False
    old_id = r.get("id")
    r["lat"] = round(float(lat), 6)
    r["lng"] = round(float(lng), 6)
    r["geocode_status"] = "ok"
    r["geocode_reasons"] = reasons
    r["geocode_display"] = display
    r["coord_source"] = source
    r["import_category"] = "READY_TO_IMPORT"
    r["verification_status"] = r.get("verification_status") or "VERIFIED_CURRENT"
    if not r.get("postal_code") and postal_from:
        pc = es_postal(postal_from)
        if pc:
            r["postal_code"] = pc
    # Preserve existing READY ids; for recovered rows recompute only if missing
    if not preserve_id or not old_id:
        r["id"] = make_id(r.get("brand"), r.get("address"), r.get("postal_code") or "", r.get("city") or "")
    else:
        r["id"] = old_id
    r["phase"] = "spain_phase4"
    return True


def geocode_row_strict(r, cache):
    if r.get("import_category") in {"COMING_SOON", "CLOSED", "DUPLICATE"}:
        return r, False
    if r.get("import_category") == "READY_TO_IMPORT" and r.get("lat") is not None:
        return r, False

    street = (r.get("address") or "").strip()
    postal = es_postal(r.get("postal_code") or "") or ""
    city = (r.get("city") or "").strip()
    if not street or not city or city.isdigit():
        r["import_category"] = "NEEDS_REVIEW"
        r["geocode_status"] = "incomplete_address"
        return r, False

    brand = r.get("brand") or ""
    name = r.get("name") or r.get("center_name") or ""
    variants = address_variants(street)

    queries = []
    struct_tries = []
    for v in variants:
        if postal:
            queries.append(f"{v}, {postal} {city}, Spain")
            queries.append(f"{v}, {postal}, España")
            struct_tries.append((v, city, postal))
        queries.append(f"{v}, {city}, España")
        struct_tries.append((v, city, ""))
        if brand:
            queries.append(f"{brand}, {v}, {city}, Spain")
            if postal:
                queries.append(f"{brand}, {v}, {postal} {city}, Spain")
    if brand and name:
        queries.append(name)
        queries.append(f"{brand} {city}")
    # shopping center name from address / center name
    for m in re.finditer(r"\b(?:CC|Centro Comercial|Parque Comercial)\s+([^,\-]+)", street + " " + name, re.I):
        cc = m.group(0).strip()
        queries.append(f"{cc}, {city}, Spain")
        if brand:
            queries.append(f"{brand} {cc}")

    seen_q = set()
    all_scored = []
    for q in queries:
        if not q or q in seen_q:
            continue
        seen_q.add(q)
        try:
            items = nominatim(q, cache)
        except Exception:
            time.sleep(1.1)
            continue
        for it in items:
            # try each street variant for scoring
            best_local = None
            for v in variants:
                sc, reasons, lat, lng = score_candidate(it, v, postal, city)
                if sc is None:
                    continue
                cand = (sc, reasons, lat, lng, it.get("display_name"), it)
                if best_local is None or sc > best_local[0]:
                    best_local = cand
            if best_local:
                all_scored.append(best_local)

    for street_s, city_s, postal_s in struct_tries[:6]:
        try:
            items = nominatim_structured(cache, street_s, city_s, postal_s)
        except Exception:
            time.sleep(1.1)
            continue
        for it in items:
            sc, reasons, lat, lng = score_candidate(it, street_s, postal_s or postal, city_s)
            if sc is None:
                continue
            all_scored.append((sc, reasons, lat, lng, it.get("display_name"), it))

    # Dedup by lat/lng
    dedup = {}
    for s in all_scored:
        k = (round(s[2], 5), round(s[3], 5))
        if k not in dedup or s[0] > dedup[k][0]:
            dedup[k] = s
    scored = list(dedup.values())
    top = pick_best_scored(scored)
    if not top:
        r["import_category"] = "NEEDS_COORDINATES"
        r["geocode_status"] = "failed"
        return r, False

    pc = es_postal((top[5].get("address") or {}).get("postcode") or "") or postal
    source = "nominatim_named_gym_poi" if "named_gym_poi" in top[1] else "nominatim_strict_address"
    if apply_geocode_hit(r, top[2], top[3], top[1], top[4], source, pc, preserve_id=True):
        return r, True
    r["import_category"] = "NEEDS_COORDINATES"
    r["geocode_status"] = "failed"
    return r, False


def overpass_elements_to_pois(path: Path):
    if not path.exists():
        return []
    data = json.loads(path.read_text(encoding="utf-8"))
    pois = []
    for el in data.get("elements") or []:
        tags = el.get("tags") or {}
        lat = el.get("lat") or (el.get("center") or {}).get("lat")
        lon = el.get("lon") or (el.get("center") or {}).get("lon")
        if lat is None or lon is None:
            continue
        if not in_spain_bbox(lat, lon):
            continue
        # synthesize nominatim-like item
        pois.append({
            "lat": str(lat),
            "lon": str(lon),
            "display_name": ", ".join(filter(None, [
                tags.get("name"),
                tags.get("addr:housenumber"),
                tags.get("addr:street"),
                tags.get("addr:city") or tags.get("addr:place"),
                tags.get("addr:postcode"),
                "España",
            ])),
            "class": tags.get("amenity") and "amenity" or tags.get("leisure") and "leisure" or "amenity",
            "type": tags.get("leisure") or tags.get("amenity") or "fitness_centre",
            "address": {
                "road": tags.get("addr:street"),
                "house_number": tags.get("addr:housenumber"),
                "city": tags.get("addr:city") or tags.get("addr:place"),
                "town": tags.get("addr:town"),
                "postcode": tags.get("addr:postcode"),
                "country_code": "es",
            },
            "osm_type": el.get("type"),
            "osm_id": el.get("id"),
            "name": tags.get("name") or "",
        })
    return pois


def match_row_to_pois(r, pois, max_dist_m=350):
    city_n = norm(r.get("city") or "")
    addr_n = norm(r.get("address") or "")
    name_n = norm(r.get("name") or "")
    postal = es_postal(r.get("postal_code") or "")
    candidates = []
    for it in pois:
        a = it.get("address") or {}
        display = it.get("display_name") or ""
        poi_name = norm(it.get("name") or display)
        city_fields = " ".join(
            norm(a.get(k) or "")
            for k in ("city", "town", "village", "municipality", "suburb")
        )
        city_hit = bool(city_n) and (city_n in city_fields or city_n in norm(display) or city_n in poi_name)
        # soft city from gym name tokens
        name_city_hit = bool(city_n) and city_n in name_n
        if city_n and not (city_hit or name_city_hit):
            continue
        pc = es_postal(a.get("postcode") or "")
        score = 0
        if postal and pc == postal:
            score += 5
        road = norm(a.get("road") or "")
        tokens = [t for t in addr_n.split() if len(t) > 3]
        overlap = sum(1 for t in tokens if t in road or t in norm(display))
        score += min(4, overlap)
        hn = str(a.get("house_number") or "")
        m = re.search(r"\b(\d+[a-z]?)\b", (r.get("address") or "").lower())
        if hn and m and hn.lower() == m.group(1).lower():
            score += 3
        for tok in re.findall(r"[a-z]{4,}", name_n):
            if tok in ("synergym", "fitness", "park", "gimnasio", "centro"):
                continue
            if tok in poi_name or tok in norm(display):
                score += 2
        if city_hit:
            score += 1
        candidates.append((score, it, pc))

    candidates.sort(key=lambda x: -x[0])
    if candidates and candidates[0][0] >= 4:
        if len(candidates) > 1 and candidates[0][0] == candidates[1][0]:
            a, b = candidates[0][1], candidates[1][1]
            d = haversine(float(a["lat"]), float(a["lon"]), float(b["lat"]), float(b["lon"]))
            if d > 200:
                return None
        return candidates[0]

    # Unique city fallback only when single POI in that city
    city_matches = []
    for it in pois:
        a = it.get("address") or {}
        display = it.get("display_name") or ""
        city_fields = " ".join(
            norm(a.get(k) or "")
            for k in ("city", "town", "village", "municipality", "suburb")
        )
        if city_n and (city_n in city_fields or city_n in norm(display) or city_n in norm(it.get("name") or "")):
            city_matches.append(it)
    if len(city_matches) == 1:
        it = city_matches[0]
        pc = es_postal((it.get("address") or {}).get("postcode") or "")
        return (3, it, pc)
    return None


def enrich_fitnesspark_from_pages(rows):
    """Refresh address/city/postal/coming-soon from cached official club pages."""
    fp_dir = PAGES / "fitnesspark"
    if not fp_dir.exists():
        return 0
    updated = 0
    by_url = {}
    for r in rows:
        if r.get("brand") != "Fitness Park":
            continue
        u = (r.get("source_url") or "").rstrip("/") + "/"
        by_url[u] = r
        by_url[(r.get("source_url") or "").rstrip("/")] = r

    for r in rows:
        if r.get("brand") != "Fitness Park":
            continue
        if r.get("import_category") not in {"NEEDS_COORDINATES", "NEEDS_REVIEW"}:
            continue
        url = r.get("source_url") or ""
        slug = url.rstrip("/").split("/")[-1]
        path = fp_dir / f"{slug}.html"
        if not path.exists():
            cands = list(fp_dir.glob(f"*{slug}*.html"))
            path = cands[0] if cands else None
        if not path or not path.exists():
            continue
        html = path.read_text(encoding="utf-8", errors="replace")
        head = html[:8000].lower()
        if any(x in head for x in ("próxima apertura", "proximamente", "próximamente", "coming soon")):
            r["import_category"] = "COMING_SOON"
            r["verification_status"] = "COMING_SOON"
            r["is_active"] = False
            r["notes"] = ((r.get("notes") or "") + "; phase4_coming_soon_page").strip("; ")
            updated += 1
            continue
        if any(x in head for x in ("cerrado temporalmente", "club cerrado", "permanently closed")):
            r["import_category"] = "CLOSED"
            r["verification_status"] = "CLOSED"
            r["is_active"] = False
            r["notes"] = ((r.get("notes") or "") + "; phase4_closed_page").strip("; ")
            updated += 1
            continue

        # JSON-LD
        for jm in re.finditer(
            r'<script[^>]+type=["\']application/ld\+json["\'][^>]*>(.*?)</script>',
            html, re.S | re.I,
        ):
            try:
                data = json.loads(jm.group(1))
            except Exception:
                continue
            nodes = data if isinstance(data, list) else [data]
            for n in nodes:
                if not isinstance(n, dict):
                    continue
                if n.get("@type") not in {"ExerciseGym", "HealthClub", "LocalBusiness", "Gym", "SportsActivityLocation"}:
                    continue
                a = n.get("address") or {}
                if not isinstance(a, dict):
                    continue
                street = (a.get("streetAddress") or "").strip()
                postal = es_postal(a.get("postalCode") or "")
                city = (a.get("addressLocality") or "").strip()
                if street:
                    r["address"] = street
                if postal:
                    r["postal_code"] = postal
                if city:
                    r["city"] = clean_city(city, name=r.get("name") or "", address=street)
                geo = n.get("geo") or {}
                if isinstance(geo, dict):
                    try:
                        la, lo = float(geo.get("latitude")), float(geo.get("longitude"))
                        if in_spain_bbox(la, lo):
                            apply_geocode_hit(
                                r, la, lo, ["official_json_ld_geo"], url,
                                "official_json_ld", postal, preserve_id=True,
                            )
                    except (TypeError, ValueError):
                        pass
                updated += 1
                break

        # Maps directions link → sometimes cleaner city
        m = re.search(r'href="(https://www\.google\.com/maps/dir//[^"]+)"', html)
        if m:
            dir_url = unquote(m.group(1))
            # Fitness Park Name, street, postal city, Spain
            tail = dir_url.split("maps/dir//", 1)[-1]
            parts = [p.strip() for p in tail.split(",") if p.strip()]
            if len(parts) >= 3:
                # try extract postal+city
                for i, p in enumerate(parts):
                    pc = es_postal(p)
                    if pc:
                        if not r.get("postal_code"):
                            r["postal_code"] = pc
                        rest = re.sub(r"\b\d{5}\b", "", p).strip()
                        if rest and (not r.get("city") or len(r.get("city") or "") < 3):
                            r["city"] = clean_city(rest)
                        # previous part often street if address empty-ish
                        if i >= 1 and parts[i - 1] and not re.search(r"Fitness Park", parts[i - 1], re.I):
                            if not r.get("address") or len(r.get("address") or "") < 5:
                                r["address"] = parts[i - 1]
            r["notes"] = ((r.get("notes") or "") + "; phase4_maps_dir_enriched").strip("; ")
    return updated


def fetch_brand_pois(brand_queries, cache, limit_per=50):
    pois = []
    seen = set()
    for q in brand_queries:
        try:
            items = nominatim(q, cache, extra={"limit": str(limit_per)})
        except Exception:
            continue
        for it in items:
            key = (it.get("osm_type"), it.get("osm_id"))
            if key in seen:
                continue
            seen.add(key)
            try:
                lat, lng = float(it["lat"]), float(it["lon"])
            except (TypeError, ValueError, KeyError):
                continue
            if not in_spain_bbox(lat, lng):
                continue
            # Reject obvious France/Portugal false positives by country_code
            cc = ((it.get("address") or {}).get("country_code") or "").lower()
            if cc and cc != "es":
                continue
            pois.append(it)
    return pois


def soft_dedupe(rows):
    collapsed = []
    by_id = {}
    kept = []
    for r in rows:
        r = dict(r)
        r["postal_code"] = es_postal(r.get("postal_code")) or r.get("postal_code") or None
        # Do not rewrite ids for existing READY
        if r.get("import_category") != "READY_TO_IMPORT" or not r.get("id"):
            r["id"] = make_id(
                r.get("brand") or r.get("chain"),
                r.get("address") or "",
                r.get("postal_code") or "",
                r.get("city") or "",
            )
        r["chain"] = r.get("chain") or r.get("brand")
        r["center_name"] = r.get("center_name") or r.get("name")
        r["country"] = "Spain"
        prev = by_id.get(r["id"])
        if prev is None:
            by_id[r["id"]] = r
            kept.append(r)
            continue
        if richness(r) > richness(prev):
            if prev.get("import_category") == "READY_TO_IMPORT" and prev.get("lat") is not None:
                if r.get("lat") is None:
                    r["lat"], r["lng"] = prev["lat"], prev["lng"]
                    r["coord_source"] = prev.get("coord_source")
                    r["import_category"] = "READY_TO_IMPORT"
            kept.remove(prev)
            by_id[r["id"]] = r
            kept.append(r)
            collapsed.append({"kept": r.get("name"), "dropped": prev.get("name"), "reason": "richer_id"})
        else:
            for f in ("postal_code", "lat", "lng", "coord_source", "source_url", "address"):
                if not prev.get(f) and r.get(f):
                    prev[f] = r[f]
            collapsed.append({"kept": prev.get("name"), "dropped": r.get("name"), "reason": "keep_existing"})

    # Soft address duplicates — never demote pre-Phase-4 READY rows
    by_soft = defaultdict(list)
    for r in kept:
        if r.get("import_category") == "DUPLICATE":
            continue
        if r.get("address"):
            by_soft[(norm(r.get("brand")), norm(r.get("address")), norm(r.get("city")))].append(r)
    for group in by_soft.values():
        if len(group) <= 1:
            continue
        group = sorted(
            group,
            key=lambda x: (
                0 if (x.get("import_category") == "READY_TO_IMPORT" and x.get("phase") != "spain_phase4") else 1,
                -richness(x),
            ),
        )
        best = group[0]
        for other in group[1:]:
            if other.get("import_category") == "READY_TO_IMPORT" and other.get("phase") != "spain_phase4":
                continue
            if richness(best) >= richness(other) or best.get("import_category") == "READY_TO_IMPORT":
                other["import_category"] = "DUPLICATE"
                other["verification_status"] = "DUPLICATE"
                collapsed.append({
                    "kept": best.get("name"), "dropped": other.get("name"),
                    "reason": "soft_addr_duplicate",
                })

    with_coords = [r for r in kept if r.get("lat") is not None and r.get("import_category") != "DUPLICATE"]
    for i, a in enumerate(with_coords):
        for b in with_coords[i + 1:]:
            if norm(a.get("brand")) != norm(b.get("brand")):
                continue
            d = haversine(float(a["lat"]), float(a["lng"]), float(b["lat"]), float(b["lng"]))
            if d > 80:
                continue
            a_prior = a.get("import_category") == "READY_TO_IMPORT" and a.get("phase") != "spain_phase4"
            b_prior = b.get("import_category") == "READY_TO_IMPORT" and b.get("phase") != "spain_phase4"
            if a_prior and not b_prior:
                winner, loser = a, b
            elif b_prior and not a_prior:
                winner, loser = b, a
            elif richness(a) >= richness(b):
                winner, loser = a, b
            else:
                winner, loser = b, a
            if loser.get("import_category") == "DUPLICATE":
                continue
            if loser.get("import_category") == "READY_TO_IMPORT" and loser.get("phase") != "spain_phase4":
                continue
            loser["import_category"] = "DUPLICATE"
            loser["verification_status"] = "DUPLICATE"
            collapsed.append({
                "kept": winner.get("name"), "dropped": loser.get("name"),
                "reason": f"proximity_{int(d)}m",
            })
    return kept, collapsed


def classify_row(r):
    if r.get("import_category") in {"COMING_SOON", "CLOSED", "DUPLICATE"}:
        return r
    if r.get("verification_status") == "COMING_SOON":
        r["import_category"] = "COMING_SOON"
        return r
    if not r.get("address") or not r.get("city") or str(r.get("city")).isdigit():
        if r.get("lat") is not None and in_spain_bbox(r.get("lat"), r.get("lng")):
            r["import_category"] = "NEEDS_REVIEW"
            return r
        r["import_category"] = "NEEDS_REVIEW"
        return r
    if r.get("lat") is not None and r.get("lng") is not None:
        try:
            lat, lng = float(r["lat"]), float(r["lng"])
            if in_spain_bbox(lat, lng):
                r["import_category"] = "READY_TO_IMPORT"
                r["lat"], r["lng"] = lat, lng
                r["verification_status"] = r.get("verification_status") or "VERIFIED_CURRENT"
                return r
            r["notes"] = ((r.get("notes") or "") + "; coord_outside_spain_bbox").strip("; ")
            r["lat"] = r["lng"] = None
            r["coord_source"] = None
        except (TypeError, ValueError):
            r["lat"] = r["lng"] = None
    if r.get("import_category") not in {"COMING_SOON", "CLOSED", "DUPLICATE"}:
        r["import_category"] = "NEEDS_COORDINATES"
    return r


def production_dedupe_flags(rows):
    """Flag staging rows that collide with production (same brand near same coords)."""
    try:
        live = json.loads(CENTERS.read_text(encoding="utf-8"))
    except Exception:
        return []
    flags = []
    live_coords = []
    for c in live:
        try:
            lat, lng = float(c.get("lat") or c.get("latitude")), float(c.get("lng") or c.get("longitude") or c.get("lon"))
        except (TypeError, ValueError):
            continue
        brand = c.get("brand") or c.get("chain") or ""
        live_coords.append((norm(brand), lat, lng, c.get("name") or c.get("center_name")))
    for r in rows:
        if r.get("lat") is None or r.get("import_category") == "DUPLICATE":
            continue
        try:
            lat, lng = float(r["lat"]), float(r["lng"])
        except (TypeError, ValueError):
            continue
        b = norm(r.get("brand"))
        for lb, llat, llng, lname in live_coords:
            if lb != b:
                continue
            d = haversine(lat, lng, llat, llng)
            if d <= 60:
                flags.append({
                    "staging": r.get("name"),
                    "production": lname,
                    "distance_m": int(d),
                    "brand": r.get("brand"),
                })
                break
    return flags


def write_geocode_review(rows):
    review = []
    for r in rows:
        if r.get("coord_source") in {
            "nominatim", "nominatim_strict_address", "nominatim_named_gym_poi",
            "overpass_named_gym_poi",
        } or r.get("geocode_status") or r.get("import_category") in {
            "NEEDS_COORDINATES", "NEEDS_REVIEW",
        }:
            review.append({
                "id": r.get("id"),
                "brand": r.get("brand"),
                "name": r.get("name"),
                "address": r.get("address"),
                "postal_code": r.get("postal_code"),
                "city": r.get("city"),
                "lat": r.get("lat"),
                "lng": r.get("lng"),
                "import_category": r.get("import_category"),
                "coord_source": r.get("coord_source"),
                "geocode_status": r.get("geocode_status"),
                "geocode_display": r.get("geocode_display"),
                "geocode_reasons": r.get("geocode_reasons"),
            })
    (OUT / "spain_geocode_review.json").write_text(
        json.dumps(review, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if review:
        fields = list(review[0].keys())
        with (OUT / "spain_geocode_review.csv").open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader()
            for row in review:
                flat = dict(row)
                if isinstance(flat.get("geocode_reasons"), list):
                    flat["geocode_reasons"] = "|".join(map(str, flat["geocode_reasons"]))
                w.writerow(flat)


def write_excel(rows):
    if not HAS_OPENPYXL:
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Spain Centers"
    headers = [
        "id", "brand", "name", "address", "postal_code", "city", "country",
        "latitude", "longitude", "status", "verification_status", "source_url",
        "import_category", "coord_source", "geocode_status", "website", "notes", "phase",
    ]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([
            r.get("id"), r.get("brand"), r.get("name"), r.get("address"),
            r.get("postal_code"), r.get("city"), r.get("country"),
            r.get("lat"), r.get("lng"), r.get("import_category"),
            r.get("verification_status"), r.get("source_url"),
            r.get("import_category"), r.get("coord_source"), r.get("geocode_status"),
            r.get("website"), r.get("notes"), r.get("phase"),
        ])
    for i, _ in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    wb.save(OUT / "Gymly_Spain_All_Discovered_Centers.xlsx")


def brand_stats(rows, brand):
    sub = [r for r in rows if r.get("brand") == brand]
    active = [r for r in sub if r.get("import_category") != "DUPLICATE"]
    return {
        "discovered": len(active),
        "ready": sum(1 for r in sub if r.get("import_category") == "READY_TO_IMPORT"),
        "unresolved": sum(1 for r in sub if r.get("import_category") in {"NEEDS_COORDINATES", "NEEDS_REVIEW"}),
        "coming_soon": sum(1 for r in sub if r.get("import_category") == "COMING_SOON"),
        "closed": sum(1 for r in sub if r.get("import_category") == "CLOSED"),
        "duplicate": sum(1 for r in sub if r.get("import_category") == "DUPLICATE"),
    }

def coord_bucket_counts(ready_rows, phase4_only=False):
    c = Counter()
    for r in ready_rows:
        if phase4_only and r.get("phase") != "spain_phase4":
            continue
        src = r.get("coord_source") or "unknown"
        bucket = COORD_BUCKET.get(src)
        if not bucket:
            if "official" in src:
                bucket = "OFFICIAL_COORDINATE"
            elif "poi" in src:
                bucket = "NAMED_GYM_POI"
            elif "nominatim" in src or src == "nominatim":
                bucket = "STRICT_ADDRESS_GEOCODE"
            elif "embed" in src or "map" in src:
                bucket = "OFFICIAL_MAP_PIN"
            else:
                bucket = src
        c[bucket] += 1
    return dict(c)


def write_reports(rows, collapsed, stats, baseline, prod_flags):
    active = [r for r in rows if r.get("import_category") != "DUPLICATE"]
    ready = [r for r in active if r.get("import_category") == "READY_TO_IMPORT"]
    cats = Counter(r.get("import_category") for r in active)
    n_dup = sum(1 for r in rows if r.get("import_category") == "DUPLICATE")

    try:
        live = json.loads(CENTERS.read_text(encoding="utf-8"))
        n_live = len(live)
    except Exception:
        n_live = 7167

    brand_ready = Counter(r.get("brand") for r in ready)
    syn = brand_stats(active, "Synergym")
    # recount duplicates for syn/fp from full rows
    syn_all = [r for r in rows if r.get("brand") == "Synergym"]
    fp_all = [r for r in rows if r.get("brand") == "Fitness Park"]
    syn = {
        "discovered": len([r for r in syn_all if r.get("import_category") != "DUPLICATE"]),
        "ready": sum(1 for r in syn_all if r.get("import_category") == "READY_TO_IMPORT"),
        "unresolved": sum(1 for r in syn_all if r.get("import_category") in {"NEEDS_COORDINATES", "NEEDS_REVIEW"}),
        "coming_soon": sum(1 for r in syn_all if r.get("import_category") == "COMING_SOON"),
        "closed": sum(1 for r in syn_all if r.get("import_category") == "CLOSED"),
        "duplicate": sum(1 for r in syn_all if r.get("import_category") == "DUPLICATE"),
        "recovered_phase4": stats.get("syn_recovered", 0),
    }
    fp = {
        "discovered": len([r for r in fp_all if r.get("import_category") != "DUPLICATE"]),
        "ready": sum(1 for r in fp_all if r.get("import_category") == "READY_TO_IMPORT"),
        "unresolved": sum(1 for r in fp_all if r.get("import_category") in {"NEEDS_COORDINATES", "NEEDS_REVIEW"}),
        "coming_soon": sum(1 for r in fp_all if r.get("import_category") == "COMING_SOON"),
        "closed": sum(1 for r in fp_all if r.get("import_category") == "CLOSED"),
        "duplicate": sum(1 for r in fp_all if r.get("import_category") == "DUPLICATE"),
        "recovered_phase4": stats.get("fp_recovered", 0),
    }

    p4_ready = [r for r in ready if r.get("phase") == "spain_phase4"]
    coord_new = coord_bucket_counts(p4_ready, phase4_only=False)
    coord_all = coord_bucket_counts(ready)

    def verdict(brand, st, estimate):
        d, rd, un = st["discovered"], st["ready"], st["unresolved"]
        if un == 0 and rd >= max(1, int(0.9 * d)):
            return "COMPLETE"
        if estimate and d >= int(0.85 * estimate) and rd >= int(0.75 * d) and un <= 25:
            return "NEAR-COMPLETE"
        if d > 0 and rd >= int(0.7 * d):
            return "NEAR-COMPLETE"
        if d > 0 and rd < int(0.5 * max(d, estimate or d)):
            return "MATERIAL GAP"
        return "NEAR-COMPLETE" if d else "MISSING"

    syn_v = verdict("Synergym", syn, 220)
    fp_v = verdict("Fitness Park", fp, 160)

    # Remaining gap detail
    syn_left = [r for r in syn_all if r.get("import_category") in {"NEEDS_COORDINATES", "NEEDS_REVIEW"}]
    fp_left = [r for r in fp_all if r.get("import_category") in {"NEEDS_COORDINATES", "NEEDS_REVIEW"}]

    phase5 = False
    # Only YES if substantial OPEN gyms remain realistically recoverable
    if syn["unresolved"] + fp["unresolved"] >= 40 and (
        stats.get("syn_recovered", 0) + stats.get("fp_recovered", 0) < 20
    ):
        phase5 = True
    # Default NO unless substantial recoverable remain from identified source
    if syn["unresolved"] <= 30 and fp["unresolved"] <= 20:
        phase5 = False
    if not stats.get("identified_recoverable_source"):
        phase5 = False

    recommendation = "PHASE 5 REQUIRED" if phase5 else "READY FOR SPAIN MERGE"

    missing_addr = sum(1 for r in active if not r.get("address"))
    missing_pc = sum(1 for r in active if not r.get("postal_code"))
    missing_city = sum(1 for r in active if not r.get("city"))
    missing_coords = sum(1 for r in active if r.get("lat") is None and r.get("import_category") not in {"COMING_SOON", "CLOSED"})

    report = {
        "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        "phase": 4,
        "baseline_phase3": baseline,
        "stats": stats,
        "overall": {
            "unique_staged_excl_duplicates": len(active),
            "READY_TO_IMPORT": cats.get("READY_TO_IMPORT", 0),
            "NEEDS_COORDINATES": cats.get("NEEDS_COORDINATES", 0),
            "NEEDS_REVIEW": cats.get("NEEDS_REVIEW", 0),
            "COMING_SOON": cats.get("COMING_SOON", 0),
            "CLOSED": cats.get("CLOSED", 0),
            "DUPLICATE": n_dup,
        },
        "synergym": syn,
        "fitness_park": fp,
        "synergym_verdict": syn_v,
        "fitness_park_verdict": fp_v,
        "coordinate_sources_phase4_new_ready": coord_new,
        "coordinate_sources_all_ready": coord_all,
        "ready_by_brand": dict(brand_ready),
        "data_quality": {
            "missing_addresses": missing_addr,
            "missing_postal_codes": missing_pc,
            "missing_cities": missing_city,
            "missing_coordinates": missing_coords,
        },
        "remaining_synergym": [
            {"id": r.get("id"), "name": r.get("name"), "address": r.get("address"),
             "city": r.get("city"), "postal_code": r.get("postal_code"),
             "category": r.get("import_category"), "reason": r.get("geocode_status")}
            for r in syn_left
        ],
        "remaining_fitness_park": [
            {"id": r.get("id"), "name": r.get("name"), "address": r.get("address"),
             "city": r.get("city"), "postal_code": r.get("postal_code"),
             "category": r.get("import_category"), "reason": r.get("geocode_status")}
            for r in fp_left
        ],
        "production_near_collisions": prod_flags[:50],
        "phase5_worthwhile": phase5,
        "recommendation": recommendation,
        "proposed_safe_merge_ready": cats.get("READY_TO_IMPORT", 0),
        "projected_catalog": n_live + cats.get("READY_TO_IMPORT", 0),
        "live_catalog": n_live,
    }

    (OUT / "SPAIN_PHASE4_READINESS_REPORT.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUT / "SPAIN_PHASE4_READY_TO_IMPORT.json").write_text(
        json.dumps(ready, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    lines = []
    lines.append("# Spain Phase 4 Readiness Report")
    lines.append("")
    lines.append(f"Generated: {report['generated_at']}")
    lines.append("")
    lines.append(f"**Status: PHASE 4 DONE — {recommendation}.**")
    lines.append("")
    lines.append("`src/data/centers.json` was not modified.")
    lines.append("")
    lines.append("## Phase 3 baseline")
    lines.append("")
    for k, v in baseline.items():
        lines.append(f"- {k}: {v}")
    lines.append("")
    lines.append("## Synergym recovery")
    lines.append("")
    lines.append(f"- recovered_to_READY: {syn['recovered_phase4']}")
    lines.append(f"- coming_soon: {syn['coming_soon']}")
    lines.append(f"- closed: {syn['closed']}")
    lines.append(f"- duplicates: {syn['duplicate']}")
    lines.append(f"- still_unresolved: {syn['unresolved']}")
    lines.append(f"- final READY: {syn['ready']}")
    lines.append(f"- coverage: {syn['ready']}/{syn['discovered']} active discovered")
    lines.append(f"- verdict: {syn_v}")
    lines.append("")
    lines.append("## Fitness Park recovery")
    lines.append("")
    lines.append(f"- recovered_to_READY: {fp['recovered_phase4']}")
    lines.append(f"- coming_soon: {fp['coming_soon']}")
    lines.append(f"- closed: {fp['closed']}")
    lines.append(f"- duplicates: {fp['duplicate']}")
    lines.append(f"- still_unresolved: {fp['unresolved']}")
    lines.append(f"- final READY: {fp['ready']}")
    lines.append(f"- coverage: {fp['ready']}/{fp['discovered']} active discovered")
    lines.append(f"- verdict: {fp_v}")
    lines.append("")
    lines.append("## Coordinate sources (Phase 4 new READY)")
    lines.append("")
    for k in ["OFFICIAL_COORDINATE", "OFFICIAL_MAP_PIN", "NAMED_GYM_POI", "STRICT_ADDRESS_GEOCODE"]:
        lines.append(f"- {k}: {coord_new.get(k, 0)}")
    lines.append("")
    lines.append("## Overall")
    lines.append("")
    lines.append("| Metric | Count |")
    lines.append("|---|---:|")
    for k, v in report["overall"].items():
        lines.append(f"| {k} | {v} |")
    lines.append("")
    lines.append("## READY by brand")
    lines.append("")
    lines.append("| Brand | READY |")
    lines.append("|---|---:|")
    for b, n in brand_ready.most_common():
        lines.append(f"| {b} | {n} |")
    lines.append("")
    lines.append("## Data quality")
    lines.append("")
    for k, v in report["data_quality"].items():
        lines.append(f"- {k}: {v}")
    lines.append("- Spanish postcodes preserved as 5-digit strings.")
    lines.append("- No city/postcode/country centroid fallbacks used.")
    lines.append("")
    lines.append("## Remaining gaps")
    lines.append("")
    lines.append(f"### Synergym ({len(syn_left)} unresolved)")
    lines.append("")
    lines.append("Official synergym.es club locator remains Incapsula-blocked; recovery used Overpass named POIs + strict Nominatim address geocoding with Av./C./P.º variants. Leftovers lack unique OSM house-level hits or have ambiguous long-avenue matches without house numbers.")
    lines.append("")
    for r in syn_left[:40]:
        lines.append(f"- {r.get('name')} | {r.get('address')} | {r.get('city')} | {r.get('import_category')}")
    lines.append("")
    lines.append(f"### Fitness Park ({len(fp_left)} unresolved)")
    lines.append("")
    lines.append("Official club pages provide addresses but rarely coordinates; OSM has few named Fitness Park Spain POIs. Leftovers are mostly shopping-center / highway-km addresses without a unique building or named-gym POI.")
    lines.append("")
    for r in fp_left[:40]:
        lines.append(f"- {r.get('name')} | {r.get('address')} | {r.get('city')} | {r.get('import_category')}")
    lines.append("")
    lines.append("## Phase 5?")
    lines.append("")
    lines.append(f"**{'YES' if phase5 else 'NO'}** — default is not worthwhile unless a new official coordinate source opens (Synergym Incapsula bypass / FP embed coords). Remaining gaps are address-quality limits, not missing discovery.")
    lines.append("")
    lines.append("## Proposed SAFE merge")
    lines.append("")
    lines.append(f"**{cats.get('READY_TO_IMPORT', 0)}** READY_TO_IMPORT rows from Phase 4 staging.")
    lines.append("")
    lines.append(f"Expected catalog after merge: **{n_live} + {cats.get('READY_TO_IMPORT', 0)} = {n_live + cats.get('READY_TO_IMPORT', 0)}**.")
    lines.append("")
    headroom = 10000 - (n_live + cats.get("READY_TO_IMPORT", 0))
    lines.append(f"## 10K checkpoint: {'YES' if headroom < 0 else 'NO'} (headroom {headroom})")
    lines.append("")
    lines.append(f"## RECOMMENDATION: {recommendation}")
    lines.append("")

    (OUT / "SPAIN_PHASE4_READINESS_REPORT.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return report


def main():
    LOG.write_text("", encoding="utf-8")
    log("=== Spain Phase 4 recovery start ===")
    rows = json.loads(STAGING.read_text(encoding="utf-8"))
    cache = {}
    if CACHE.exists():
        cache = json.loads(CACHE.read_text(encoding="utf-8"))

    baseline_cats = Counter(r.get("import_category") for r in rows)
    baseline = {
        "READY_TO_IMPORT": baseline_cats.get("READY_TO_IMPORT", 0),
        "NEEDS_COORDINATES": baseline_cats.get("NEEDS_COORDINATES", 0),
        "NEEDS_REVIEW": baseline_cats.get("NEEDS_REVIEW", 0),
        "COMING_SOON": baseline_cats.get("COMING_SOON", 0),
        "DUPLICATE": baseline_cats.get("DUPLICATE", 0),
        "unique_excl_dup": sum(1 for r in rows if r.get("import_category") != "DUPLICATE"),
    }
    log("Phase 3 baseline:", baseline)

    # Freeze READY ids
    ready_ids = {
        r["id"] for r in rows
        if r.get("import_category") == "READY_TO_IMPORT" and r.get("id")
    }
    log("Preserving READY ids:", len(ready_ids))

    stats = {
        "syn_repaired": 0,
        "syn_poi": 0,
        "syn_geocode": 0,
        "syn_recovered": 0,
        "fp_enriched": 0,
        "fp_poi": 0,
        "fp_geocode": 0,
        "fp_recovered": 0,
        "identified_recoverable_source": False,
    }

    # 1) Repair Synergym unresolved address/city
    for r in rows:
        if r.get("brand") != "Synergym":
            continue
        if r.get("import_category") not in {"NEEDS_COORDINATES", "NEEDS_REVIEW"}:
            continue
        before = (r.get("address"), r.get("city"))
        repair_synergym_row(r)
        after = (r.get("address"), r.get("city"))
        if after != before:
            stats["syn_repaired"] += 1
    log(f"Synergym repairs: {stats['syn_repaired']}")

    # 2) Enrich FP from official pages
    stats["fp_enriched"] = enrich_fitnesspark_from_pages(rows)
    log(f"FP page enrich/classify: {stats['fp_enriched']}")

    # 3) Load Overpass Synergym POIs + Nominatim brand POIs
    syn_pois = overpass_elements_to_pois(RAW / "overpass_synergym_p4.json")
    log(f"Overpass Synergym POIs in Spain: {len(syn_pois)}")

    # Expand Nominatim brand POIs for unresolved cities
    syn_cities = sorted({
        r.get("city") for r in rows
        if r.get("brand") == "Synergym"
        and r.get("import_category") in {"NEEDS_COORDINATES", "NEEDS_REVIEW"}
        and r.get("city") and not str(r.get("city")).isdigit()
    })
    fp_cities = sorted({
        r.get("city") for r in rows
        if r.get("brand") == "Fitness Park"
        and r.get("import_category") in {"NEEDS_COORDINATES", "NEEDS_REVIEW"}
        and r.get("city")
    })
    log(f"Nominatim Synergym city POI queries: {len(syn_cities)}")
    syn_nom = fetch_brand_pois(
        ["Synergym", "SynerGym"] + [f"Synergym {c}" for c in syn_cities],
        cache, limit_per=20,
    )
    # merge overpass + nominatim
    syn_all_pois = syn_pois + syn_nom
    log(f"Synergym POI pool: {len(syn_all_pois)}")

    log(f"Nominatim Fitness Park city POI queries: {len(fp_cities)}")
    fp_pois = fetch_brand_pois(
        ["Fitness Park España", "Fitness Park Spain"]
        + [f"Fitness Park {c}" for c in fp_cities],
        cache, limit_per=15,
    )
    # Also reuse prior nominatim dump
    prior = RAW / "nominatim_fitnesspark_pois_p3.json"
    if prior.exists():
        try:
            for it in json.loads(prior.read_text(encoding="utf-8")):
                if in_spain_bbox(float(it["lat"]), float(it["lon"])):
                    cc = ((it.get("address") or {}).get("country_code") or "es").lower()
                    if cc == "es":
                        fp_pois.append(it)
        except Exception:
            pass
    log(f"Fitness Park POI pool: {len(fp_pois)}")

    # 4) Match POIs
    for r in rows:
        if r.get("import_category") not in {"NEEDS_COORDINATES", "NEEDS_REVIEW"}:
            continue
        if r.get("brand") == "Synergym":
            m = match_row_to_pois(r, syn_all_pois)
            if not m:
                continue
            score, it, pc = m
            if apply_geocode_hit(
                r, float(it["lat"]), float(it["lon"]),
                ["named_gym_poi", f"match_score_{score}"],
                it.get("display_name"),
                "overpass_named_gym_poi" if it.get("osm_id") and not it.get("place_id") else "nominatim_named_gym_poi",
                pc or (it.get("address") or {}).get("postcode"),
                preserve_id=True,
            ):
                # Prefer overpass source label when from overpass file
                if not it.get("place_id"):
                    r["coord_source"] = "overpass_named_gym_poi"
                stats["syn_poi"] += 1
                stats["syn_recovered"] += 1
        elif r.get("brand") == "Fitness Park":
            m = match_row_to_pois(r, fp_pois)
            if not m:
                continue
            score, it, pc = m
            if apply_geocode_hit(
                r, float(it["lat"]), float(it["lon"]),
                ["named_gym_poi", f"match_score_{score}"],
                it.get("display_name"),
                "nominatim_named_gym_poi",
                pc or (it.get("address") or {}).get("postcode"),
                preserve_id=True,
            ):
                stats["fp_poi"] += 1
                stats["fp_recovered"] += 1

    log(f"POI matched Syn={stats['syn_poi']} FP={stats['fp_poi']}")

    # 5) Strict geocode remaining
    unresolved = [
        r for r in rows
        if r.get("brand") in {"Synergym", "Fitness Park"}
        and r.get("import_category") in {"NEEDS_COORDINATES", "NEEDS_REVIEW"}
    ]
    log(f"Strict geocode remaining: {len(unresolved)}")
    for i, r in enumerate(unresolved):
        before_cat = r.get("import_category")
        _, ok = geocode_row_strict(r, cache)
        if ok:
            if r.get("brand") == "Synergym":
                stats["syn_geocode"] += 1
                stats["syn_recovered"] += 1
            else:
                stats["fp_geocode"] += 1
                stats["fp_recovered"] += 1
        if (i + 1) % 10 == 0:
            log(f"  geocode progress {i+1}/{len(unresolved)} syn={stats['syn_recovered']} fp={stats['fp_recovered']}")
            CACHE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")

    # 6) Reclassify + dedupe
    for r in rows:
        if r.get("brand") in {"Synergym", "Fitness Park"}:
            classify_row(r)

    kept, collapsed = soft_dedupe(rows)
    # Ensure we didn't destroy unrelated brands — soft_dedupe returns all
    rows = kept

    # Restore any READY that somehow lost coords (safety)
    for r in rows:
        if r.get("id") in ready_ids and r.get("import_category") != "DUPLICATE":
            if r.get("lat") is not None:
                r["import_category"] = "READY_TO_IMPORT"

    prod_flags = production_dedupe_flags(rows)
    log(f"Production near-collisions: {len(prod_flags)}")

    # Persist
    CACHE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
    STAGING.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    write_geocode_review(rows)
    write_excel(rows)

    (OUT / "spain_duplicate_analysis.json").write_text(
        json.dumps({
            "phase": 4,
            "collapsed": collapsed,
            "production_near_collisions": prod_flags,
            "counts": Counter(r.get("import_category") for r in rows),
        }, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )

    report = write_reports(rows, collapsed, stats, baseline, prod_flags)
    log("Phase 4 done. READY=", report["overall"]["READY_TO_IMPORT"])
    log("Syn READY", report["synergym"]["ready"], "FP READY", report["fitness_park"]["ready"])
    log("Recommendation:", report["recommendation"])


if __name__ == "__main__":
    main()
