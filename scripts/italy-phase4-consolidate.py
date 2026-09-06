#!/usr/bin/env python3
"""
Italy Phase 4 consolidate — geocode gap recoveries, preserve Phase 3 READY,
dedupe, validate, and emit completeness import set.

Does NOT modify src/data/centers.json.
Projected catalog = 9056 + FINAL_NEW_READY (do not re-add 550).
"""
from __future__ import annotations

import hashlib
import json
import math
import re
import ssl
import time
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data/italy"
SCRAPES = OUT / "scrapes"
CENTERS = ROOT / "src/data/centers.json"
STAGING = OUT / "italy_centers_staging.json"
CACHE = OUT / "italy_geocode_cache.json"
P3_CACHE = OUT / "italy_phase3_geocode_cache.json"
P4_CACHE = OUT / "italy_phase4_geocode_cache.json"
P3_READY = OUT / "ITALY_PHASE3_READY_TO_IMPORT.json"
P3_CANDIDATES = OUT / "ITALY_PHASE3_NEW_CANDIDATES.json"
DISCOVERY = SCRAPES / "phase4_discovery.json"
DISCOVERY_META = SCRAPES / "phase4_discovery_meta.json"

ctx = ssl.create_default_context()
UA = "GymlyItalyGeocoder/4.0 (catalog research; accuracy over coverage; no fallback centroids)"

IT_MAINLAND = (36.6, 47.15, 6.6, 18.6)
IT_SICILY = (36.6, 38.35, 12.0, 15.7)
IT_SARDINIA = (38.8, 41.35, 8.1, 9.9)

ALLOWED_COORD = {
    "OFFICIAL_COORDINATE",
    "OFFICIAL_MAP_PIN",
    "NAMED_GYM_POI",
    "STRICT_ADDRESS_GEOCODE",
}
COARSE = {
    "country",
    "state",
    "region",
    "county",
    "postcode",
    "municipality",
    "city",
    "town",
    "village",
    "suburb",
    "neighbourhood",
    "quarter",
    "district",
    "borough",
    "province",
    "island",
}
FORBIDDEN_CC = {"sm", "va", "fr", "ch", "at", "si", "hr", "mt"}


def log(*a):
    print(*a, flush=True)


def make_id(brand, address, postal, city):
    key = "|".join(
        [
            (brand or "").strip().lower(),
            (address or "").strip().lower(),
            (str(postal) or "").strip().lower(),
            (city or "").strip().lower(),
            "italy",
        ]
    )
    return "it_" + hashlib.md5(key.encode("utf-8")).hexdigest()[:10]


def norm(s):
    s = (s or "").lower()
    for a, b in [
        ("à", "a"),
        ("á", "a"),
        ("è", "e"),
        ("é", "e"),
        ("ì", "i"),
        ("í", "i"),
        ("ò", "o"),
        ("ó", "o"),
        ("ù", "u"),
        ("ú", "u"),
        ("’", "'"),
        ("'", " "),
    ]:
        s = s.replace(a, b)
    s = re.sub(r"\bv\.le\b", "viale", s)
    s = re.sub(r"\bc\.so\b", "corso", s)
    s = re.sub(r"\bp\.zza\b", "piazza", s)
    s = re.sub(r"\bp\.le\b", "piazzale", s)
    s = re.sub(r"\bstr\.\b", "strada", s)
    s = re.sub(r"\bs\.\s*", "san ", s)
    s = re.sub(r"\bmastri\b", "maestri", s)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def haversine(lat1, lng1, lat2, lng2):
    R = 6371000
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lng2 - lng1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * R * math.asin(math.sqrt(a))


def it_postal(s) -> str:
    if s is None:
        return ""
    if isinstance(s, float) and math.isnan(s):
        return ""
    m = re.search(r"\b(\d{5})\b", str(s).strip())
    return m.group(1) if m else ""


def valid_city(city) -> str:
    city = str(city or "").strip(" ,.")
    if not city or city.isdigit() or it_postal(city) == city:
        return ""
    if len(city) == 2 and city.isupper():
        return ""
    return re.sub(r"\s+[A-Z]{2}$", "", city).strip()


def in_italy_bbox(lat, lng) -> bool:
    try:
        lat, lng = float(lat), float(lng)
    except (TypeError, ValueError):
        return False
    if not (math.isfinite(lat) and math.isfinite(lng)) or (lat == 0 and lng == 0):
        return False
    for lo, hi, w, e in (IT_MAINLAND, IT_SICILY, IT_SARDINIA):
        if lo <= lat <= hi and w <= lng <= e:
            return True
    return False


def load_cache():
    cache = {}
    for path in (CACHE, P3_CACHE, P4_CACHE):
        if path.exists():
            try:
                cache.update(json.loads(path.read_text(encoding="utf-8")))
            except Exception:
                pass
    return cache


def save_cache(cache):
    P4_CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")


def nominatim(query, cache):
    if query in cache:
        return cache[query]
    url = "https://nominatim.openstreetmap.org/search?" + urllib.parse.urlencode(
        {
            "q": query,
            "format": "json",
            "addressdetails": 1,
            "limit": 5,
            "countrycodes": "it",
        }
    )
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, context=ctx, timeout=30) as r:
        data = json.loads(r.read().decode())
    cache[query] = data
    time.sleep(1.1)
    return data


def nominatim_reverse(lat, lng, cache):
    key = f"rev:{round(float(lat), 6)},{round(float(lng), 6)}"
    if key in cache:
        return cache[key]
    url = "https://nominatim.openstreetmap.org/reverse?" + urllib.parse.urlencode(
        {
            "lat": lat,
            "lon": lng,
            "format": "json",
            "addressdetails": 1,
            "zoom": 18,
        }
    )
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, context=ctx, timeout=30) as r:
        data = json.loads(r.read().decode())
    cache[key] = data
    time.sleep(1.1)
    return data


def street_tokens(street):
    stop = {
        "via",
        "viale",
        "corso",
        "piazza",
        "piazzale",
        "strada",
        "contrada",
        "localita",
        "loc",
        "vicolo",
        "largo",
        "dei",
        "degli",
        "delle",
        "del",
        "della",
        "di",
        "da",
        "san",
        "santa",
        "sant",
        "don",
        "centro",
        "commerciale",
        "cc",
    }
    return [t for t in norm(street).split() if len(t) > 2 and t not in stop]


def score_candidate(item, street, postal, city, brand):
    addr = item.get("address") or {}
    lat, lng = float(item["lat"]), float(item["lon"])
    if not in_italy_bbox(lat, lng):
        return None, ["outside_italy"], lat, lng
    cc = (addr.get("country_code") or "").lower()
    if cc and cc != "it":
        return None, ["not_country_it:" + cc], lat, lng
    if cc in FORBIDDEN_CC:
        return None, ["forbidden_cc:" + cc], lat, lng
    t = (item.get("type") or "").lower()
    addresstype = (item.get("addresstype") or "").lower()
    osm_class = (item.get("class") or "").lower()
    if t in COARSE or addresstype in COARSE:
        return None, ["coarse_type:" + (addresstype or t)], lat, lng
    if osm_class in {"boundary", "place"} and t not in {
        "house",
        "building",
        "yes",
        "retail",
        "commercial",
        "industrial",
        "gym",
        "fitness_centre",
        "sports_centre",
        "sports_hall",
    }:
        return None, ["coarse_class:" + osm_class + ":" + t], lat, lng
    if osm_class == "highway":
        return None, ["highway_centroid"], lat, lng

    display = (item.get("display_name") or "").lower()
    name = (item.get("name") or "").lower()
    if "uffici icon" in display:
        return None, ["icon_hq_offices"], lat, lng

    score = 0
    reasons = []
    pc = re.sub(r"\s+", "", str(addr.get("postcode") or "")).upper()
    postal_n = re.sub(r"\s+", "", str(postal or "")).upper()
    postal_ok = bool(postal_n and pc == postal_n)
    if postal_ok:
        score += 5
        reasons.append("postal_exact")
    elif postal_n and pc and pc[:2] == postal_n[:2]:
        score += 1
        reasons.append("postal_soft")

    city_n = norm(city)
    city_fields = " ".join(
        norm(addr.get(k) or "")
        for k in ("city", "town", "village", "municipality", "suburb", "city_district")
    )
    if city_n and (city_n in city_fields or city_n in norm(display)):
        score += 3
        reasons.append("city_ok")

    toks = street_tokens(street)
    road = norm(addr.get("road") or "")
    road_ok = False
    if toks and (
        any(t in road for t in toks)
        or any(t in norm(display) for t in toks)
        or (road and any(t in road for t in toks))
    ):
        score += 4
        reasons.append("road_match")
        road_ok = True

    hn = str(addr.get("house_number") or "")
    m = re.search(r"\b(\d+)\b", street or "")
    hn_ok = False
    if hn and m and hn.split("/")[0] == m.group(1):
        score += 3
        reasons.append("house_number_match")
        hn_ok = True

    if osm_class in {"building", "amenity", "leisure", "shop", "office"} or t in {
        "gym",
        "fitness_centre",
        "sports_centre",
        "sports_hall",
        "yes",
        "retail",
        "commercial",
        "mall",
    }:
        score += 2
        reasons.append("building_or_amenity")

    brand_n = norm(brand).replace(" ", "")
    blob = norm(display) + " " + norm(name)
    named = bool(brand_n and brand_n in blob.replace(" ", ""))
    if any(
        x in display
        for x in (
            "fit express",
            "fitexpress",
            "icon palestre",
            "iconpalestre",
            "fitness park",
            "fitinn",
        )
    ):
        named = True
    if named:
        score += 5
        reasons.append("named_gym_poi")

    # Named gym without street tokens OR exact postal → reject (avoids wrong-club city hits)
    if named and not road_ok and not postal_ok and not hn_ok:
        return None, reasons + ["named_without_street_or_postal"], lat, lng
    if not named and not road_ok and not hn_ok:
        return None, reasons + ["no_street_or_hn"], lat, lng
    if score < 7:
        return None, reasons + ["score_too_low"], lat, lng

    return score, reasons, lat, lng


def fill_cap_from_reverse(r, cache):
    if it_postal(r.get("postal_code") or ""):
        return r
    if r.get("coord_source") not in {"OFFICIAL_COORDINATE", "OFFICIAL_MAP_PIN", "NAMED_GYM_POI"}:
        return r
    if r.get("lat") is None or r.get("lng") is None:
        return r
    try:
        data = nominatim_reverse(r["lat"], r["lng"], cache)
    except Exception:
        time.sleep(1.1)
        return r
    addr = (data or {}).get("address") or {}
    cc = (addr.get("country_code") or "").lower()
    if cc != "it" or cc in FORBIDDEN_CC:
        return r
    postal = it_postal(addr.get("postcode") or "")
    if not postal:
        return r
    old_id = r["id"]
    r["postal_code"] = postal
    if not valid_city(r.get("city") or ""):
        city = (
            addr.get("city")
            or addr.get("town")
            or addr.get("village")
            or addr.get("municipality")
        )
        r["city"] = valid_city(city or "")
    r["notes"] = ((r.get("notes") or "") + "; cap_from_pin_reverse").strip("; ")
    r["id"] = make_id(r.get("brand"), r.get("address") or "", postal, r.get("city") or "")
    if r["id"] != old_id:
        r["notes"] = ((r.get("notes") or "") + f"; id_rekeyed_from_{old_id}").strip("; ")
    return r


def geocode_row(r, cache, review):
    if r.get("import_category") in {"COMING_SOON", "CLOSED", "DUPLICATE", "ALREADY_LIVE"}:
        return r

    # Already have allowed coords + CAP
    if (
        r.get("lat") is not None
        and r.get("lng") is not None
        and r.get("coord_source") in ALLOWED_COORD
        and it_postal(r.get("postal_code") or "")
        and in_italy_bbox(r["lat"], r["lng"])
    ):
        r["import_category"] = "READY_TO_IMPORT"
        r["is_active"] = True
        return r

    # Official / named pin missing CAP → reverse
    if (
        r.get("lat") is not None
        and r.get("lng") is not None
        and r.get("coord_source") in {"OFFICIAL_COORDINATE", "OFFICIAL_MAP_PIN", "NAMED_GYM_POI"}
    ):
        r = fill_cap_from_reverse(r, cache)
        if it_postal(r.get("postal_code") or "") and in_italy_bbox(r["lat"], r["lng"]):
            r["import_category"] = "READY_TO_IMPORT"
            r["is_active"] = True
            return r
        r["import_category"] = "NEEDS_REVIEW"
        r["notes"] = ((r.get("notes") or "") + "; missing_cap_after_reverse").strip("; ")
        return r

    street = r.get("address") or ""
    postal = it_postal(r.get("postal_code") or "")
    city = valid_city(r.get("city") or "")
    # Strip leaked email crumbs / province tails from Phase 4 Icon parses
    city = re.sub(r"\s+[a-z][a-z0-9]{2,}$", "", city).strip()
    city = re.sub(r"\s+[A-Z]{2}\b.*$", "", city).strip()
    city = valid_city(city)
    brand = r.get("brand") or ""
    name = r.get("name") or ""
    r["city"] = city

    if not street or not city:
        r["import_category"] = "NEEDS_REVIEW"
        r["geocode_status"] = "incomplete_address"
        return r

    queries = [
        f"{brand}, {street}, {postal} {city}, Italy".strip(", "),
        f"{brand} {street} {city}",
        f"{name}, {city}, Italy",
        f"{street}, {postal} {city}, Italy",
        f"{street}, {city}, Italy",
    ]

    best = None
    best_score = -1
    best_disp = None
    tried = []
    for q in queries:
        try:
            results = nominatim(q, cache)
        except Exception as e:
            tried.append({"q": q, "error": str(e)})
            time.sleep(1.1)
            continue
        tried.append({"q": q, "n": len(results or [])})
        for item in results or []:
            sc, reasons, lat, lng = score_candidate(item, street, postal, city, brand)
            if sc is None:
                continue
            src = "NAMED_GYM_POI" if "named_gym_poi" in reasons else "STRICT_ADDRESS_GEOCODE"
            if sc > best_score:
                best_score = sc
                best = (lat, lng, src, reasons)
                best_disp = item.get("display_name")

    if best:
        lat, lng, src, reasons = best
        r["lat"], r["lng"] = lat, lng
        r["coord_source"] = src
        r["notes"] = (
            (r.get("notes") or "") + f"; geocode={src}:{','.join(reasons)}"
        ).strip("; ")
        if not it_postal(r.get("postal_code") or ""):
            try:
                data = nominatim_reverse(lat, lng, cache)
                addr = (data or {}).get("address") or {}
                pc = it_postal(addr.get("postcode") or "")
                if pc and (addr.get("country_code") or "").lower() == "it":
                    r["postal_code"] = pc
                    r["id"] = make_id(
                        r.get("brand"), r.get("address") or "", pc, r.get("city") or ""
                    )
                    r["notes"] = (
                        (r.get("notes") or "") + "; cap_from_geocode_reverse"
                    ).strip("; ")
            except Exception:
                time.sleep(1.1)
        if it_postal(r.get("postal_code") or "") and in_italy_bbox(lat, lng):
            r["import_category"] = "READY_TO_IMPORT"
            r["is_active"] = True
            r["geocode_status"] = "ok"
        else:
            r["import_category"] = "NEEDS_REVIEW"
            r["geocode_status"] = "missing_cap"
        review.append(
            {
                "id": r["id"],
                "name": r["name"],
                "status": r["import_category"],
                "tried": tried,
                "best": best_disp,
            }
        )
        return r

    r["import_category"] = "NEEDS_COORDINATES" if postal else "NEEDS_REVIEW"
    r["geocode_status"] = "no_high_confidence_hit"
    review.append(
        {"id": r.get("id"), "name": r.get("name"), "status": r["import_category"], "tried": tried}
    )
    return r


def ready_gate(r) -> list[str]:
    errs = []
    if not (r.get("id") or "").startswith("it_"):
        errs.append("bad_id")
    if not r.get("name"):
        errs.append("missing_name")
    if not r.get("brand"):
        errs.append("missing_brand")
    if r.get("brand") == "Icon" or (
        r.get("brand") and r.get("brand").lower() == "icon" and "palestre" not in r.get("brand").lower()
    ):
        errs.append("generic_icon_brand")
    if not r.get("address"):
        errs.append("missing_address")
    if not re.fullmatch(r"\d{5}", str(r.get("postal_code") or "")):
        errs.append("bad_cap")
    if not valid_city(r.get("city") or ""):
        errs.append("bad_city")
    if (r.get("country") or "") != "Italy":
        errs.append("bad_country")
    try:
        lat, lng = float(r["lat"]), float(r["lng"])
        if not math.isfinite(lat) or not math.isfinite(lng):
            errs.append("nonfinite_coords")
        elif not in_italy_bbox(lat, lng):
            errs.append("coords_outside_italy")
    except (TypeError, ValueError, KeyError):
        errs.append("missing_coords")
    if r.get("coord_source") not in ALLOWED_COORD:
        errs.append("bad_coord_source")
    if r.get("import_category") in {"CLOSED", "COMING_SOON"}:
        errs.append("not_open")
    # mojibake / foreign
    blob = " ".join(
        str(r.get(k) or "") for k in ("name", "address", "city", "brand")
    )
    if re.search(r"Ã.|Â.|â€", blob):
        errs.append("mojibake")
    return errs


def dedupe_analysis(candidates, production_all, production_it, phase3_unresolved):
    report = {
        "same_id_vs_italy": [],
        "same_id_vs_catalog": [],
        "same_id_vs_phase3_ready": [],
        "same_brand_proximity_50m": [],
        "same_brand_proximity_100m": [],
        "cross_brand_proximity_50m": [],
        "internal_dupes": [],
        "vs_phase3_unresolved_notes": [],
    }
    prod_ids = {c["id"] for c in production_all}
    it_ids = {c["id"] for c in production_it}
    by_brand = defaultdict(list)
    for c in production_it:
        if c.get("lat") is not None and c.get("lng") is not None:
            by_brand[norm(c.get("brand"))].append(c)
    all_with_coords = [
        c for c in production_all if c.get("lat") is not None and c.get("lng") is not None
    ]

    seen = {}
    kept = []
    for r in candidates:
        if r["id"] in it_ids:
            report["same_id_vs_italy"].append(r["id"])
            r["import_category"] = "ALREADY_LIVE"
            continue
        if r["id"] in prod_ids:
            report["same_id_vs_catalog"].append(r["id"])
            r["import_category"] = "DUPLICATE"
            continue
        if r["id"] in seen:
            report["internal_dupes"].append(r["id"])
            continue
        seen[r["id"]] = r

        if r.get("lat") is not None and r.get("lng") is not None and r.get("import_category") != "CLOSED":
            try:
                la, lo = float(r["lat"]), float(r["lng"])
            except (TypeError, ValueError):
                kept.append(r)
                continue
            brand_n = norm(r.get("brand"))
            dup = False
            for p in by_brand.get(brand_n, []):
                d = haversine(la, lo, float(p["lat"]), float(p["lng"]))
                if d < 50:
                    report["same_brand_proximity_50m"].append(
                        {
                            "new": r["name"],
                            "live": p.get("name"),
                            "m": round(d, 1),
                            "new_id": r["id"],
                            "live_id": p["id"],
                        }
                    )
                    r["import_category"] = "DUPLICATE"
                    r["notes"] = (
                        (r.get("notes") or "") + f"; dup_same_brand_{d:.0f}m"
                    ).strip("; ")
                    dup = True
                    break
                if d < 100:
                    report["same_brand_proximity_100m"].append(
                        {
                            "new": r["name"],
                            "live": p.get("name"),
                            "m": round(d, 1),
                            "new_id": r["id"],
                            "live_id": p["id"],
                        }
                    )
            if dup:
                continue
            for p in all_with_coords:
                if norm(p.get("brand")) == brand_n:
                    continue
                try:
                    d = haversine(la, lo, float(p["lat"]), float(p["lng"]))
                except (TypeError, ValueError):
                    continue
                if d < 50:
                    report["cross_brand_proximity_50m"].append(
                        {
                            "new": r["name"],
                            "live": p.get("name"),
                            "live_brand": p.get("brand"),
                            "m": round(d, 1),
                        }
                    )
            # also vs other new candidates already kept
            for p in kept:
                if p.get("lat") is None or norm(p.get("brand")) != brand_n:
                    continue
                d = haversine(la, lo, float(p["lat"]), float(p["lng"]))
                if d < 50:
                    report["same_brand_proximity_50m"].append(
                        {
                            "new": r["name"],
                            "live": p.get("name"),
                            "m": round(d, 1),
                            "new_id": r["id"],
                            "live_id": p["id"],
                            "scope": "internal_phase4",
                        }
                    )
                    r["import_category"] = "DUPLICATE"
                    dup = True
                    break
            if dup:
                continue
        kept.append(r)
    return kept, report


def validate_final(rows):
    issues = {
        "dup_ids": [],
        "already_in_production": [],
        "bad_cap": [],
        "bad_address": [],
        "bad_city": [],
        "bad_coords": [],
        "bad_fallback": [],
        "foreign": [],
        "mojibake": [],
        "closed_or_coming": [],
        "gate_failures": [],
    }
    prod_ids = {c["id"] for c in json.loads(CENTERS.read_text(encoding="utf-8"))}
    seen = set()
    for r in rows:
        if r["id"] in seen:
            issues["dup_ids"].append(r["id"])
        seen.add(r["id"])
        if r["id"] in prod_ids:
            issues["already_in_production"].append(r["id"])
        errs = ready_gate(r)
        if errs:
            issues["gate_failures"].append({"id": r["id"], "name": r["name"], "errs": errs})
        if "bad_cap" in errs:
            issues["bad_cap"].append(r["id"])
        if "missing_address" in errs:
            issues["bad_address"].append(r["id"])
        if "bad_city" in errs:
            issues["bad_city"].append(r["id"])
        if any(x in errs for x in ("missing_coords", "nonfinite_coords", "coords_outside_italy", "bad_coord_source")):
            issues["bad_coords"].append(r["id"])
        if "mojibake" in errs:
            issues["mojibake"].append(r["id"])
        if r.get("import_category") in {"CLOSED", "COMING_SOON"}:
            issues["closed_or_coming"].append(r["id"])
    issues["ok"] = all(len(v) == 0 for k, v in issues.items() if k != "ok")
    return issues


def write_xlsx(phase3_ready, phase4_new, final_ready, unresolved, meta):
    if not HAS_OPENPYXL:
        log("openpyxl missing — skip xlsx")
        return
    wb = Workbook()

    def sheet(title, rows, headers):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(h) if h != "postal_code" else str(r.get("postal_code") or "") for h in headers])
            # force CAP text
            if "postal_code" in headers:
                col = headers.index("postal_code") + 1
                cell = ws.cell(ws.max_row, col)
                cell.number_format = "@"
                cell.value = str(r.get("postal_code") or "")
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(40, max(12, len(h) + 2))
        return ws

    # remove default
    wb.remove(wb.active)
    headers = [
        "id",
        "discovery_class",
        "import_category",
        "brand",
        "name",
        "address",
        "postal_code",
        "city",
        "country",
        "lat",
        "lng",
        "coord_source",
        "source_url",
        "notes",
        "phase",
    ]
    sheet("FINAL_READY", final_ready, headers)
    sheet("Phase3_READY_retained", phase3_ready, headers)
    sheet("Phase4_NEW_recovered", phase4_new, headers)
    sheet("Unresolved", unresolved, headers)

    ws = wb.create_sheet("Summary")
    for k, v in meta.items():
        if isinstance(v, (dict, list)):
            ws.append([k, json.dumps(v, ensure_ascii=False)[:500]])
        else:
            ws.append([k, v])

    path = OUT / "Gymly_Italy_Final_Completeness.xlsx"
    wb.save(path)
    log("Wrote", path)


def main():
    log("=== Italy Phase 4 consolidate ===")
    production_all = json.loads(CENTERS.read_text(encoding="utf-8"))
    production_it = [
        c for c in production_all if (c.get("country") or "").lower() == "italy"
    ]
    assert len(production_all) == 9056, f"catalog size {len(production_all)}"
    assert len(production_it) == 550, f"italy size {len(production_it)}"

    phase3_ready = json.loads(P3_READY.read_text(encoding="utf-8"))
    assert len(phase3_ready) == 22, f"phase3 ready {len(phase3_ready)}"
    phase3_cands = json.loads(P3_CANDIDATES.read_text(encoding="utf-8"))
    phase3_unresolved = [
        r
        for r in phase3_cands
        if r.get("import_category") not in {"READY_TO_IMPORT"}
    ]

    discovery = json.loads(DISCOVERY.read_text(encoding="utf-8"))
    discovery_meta = (
        json.loads(DISCOVERY_META.read_text(encoding="utf-8"))
        if DISCOVERY_META.exists()
        else {}
    )

    cache = load_cache()
    review = []

    # Geocode discovery rows
    recovered = []
    for r in discovery:
        rr = deepcopy(r)
        rr = geocode_row(rr, cache, review)
        recovered.append(rr)
    save_cache(cache)

    # Phase 3 READY retention / demotion check
    demotions = []
    retained_p3 = []
    for r in phase3_ready:
        rr = deepcopy(r)
        errs = ready_gate(rr)
        # Re-verify not already live / not closed
        if rr.get("verification_status") == "CLOSED" or rr.get("import_category") == "CLOSED":
            demotions.append({"id": rr["id"], "name": rr["name"], "reason": "CLOSED"})
            continue
        if rr["id"] in {c["id"] for c in production_all}:
            demotions.append(
                {"id": rr["id"], "name": rr["name"], "reason": "ALREADY_IN_PRODUCTION"}
            )
            continue
        if errs:
            demotions.append(
                {"id": rr["id"], "name": rr["name"], "reason": "READY_GATE:" + ",".join(errs)}
            )
            continue
        rr["discovery_class"] = rr.get("discovery_class") or "PHASE3_READY_RETAINED"
        rr["phase"] = rr.get("phase") or "italy_phase3"
        retained_p3.append(rr)

    log(f"Phase 3 READY retained {len(retained_p3)} / demoted {len(demotions)}")
    for d in demotions:
        log("  DEMOTE", d)

    # Phase 4 newly recovered READY (not already in Phase 3 READY ids)
    p3_ids = {r["id"] for r in retained_p3}
    phase4_ready_raw = [
        r
        for r in recovered
        if r.get("import_category") == "READY_TO_IMPORT" and r["id"] not in p3_ids
    ]

    # Dedupe Phase 4 new vs production + phase3
    phase4_kept, dup_report = dedupe_analysis(
        phase4_ready_raw, production_all, production_it, phase3_unresolved
    )
    # Also flag id collision with phase3 ready
    phase4_new = []
    for r in phase4_kept:
        if r.get("import_category") != "READY_TO_IMPORT":
            continue
        if r["id"] in p3_ids:
            dup_report["same_id_vs_phase3_ready"].append(r["id"])
            continue
        # final gate
        errs = ready_gate(r)
        if errs:
            r["import_category"] = "NEEDS_REVIEW"
            r["notes"] = ((r.get("notes") or "") + "; gate:" + ",".join(errs)).strip("; ")
            continue
        r["discovery_class"] = r.get("discovery_class") or "PHASE4_RECOVERED"
        phase4_new.append(r)

    # Same-brand proximity inspect between phase4_new and retained_p3
    for r in list(phase4_new):
        if r.get("lat") is None:
            continue
        for p in retained_p3:
            if norm(p.get("brand")) != norm(r.get("brand")):
                continue
            if p.get("lat") is None:
                continue
            d = haversine(float(r["lat"]), float(r["lng"]), float(p["lat"]), float(p["lng"]))
            if d < 50:
                dup_report["same_brand_proximity_50m"].append(
                    {
                        "new": r["name"],
                        "live": p["name"],
                        "m": round(d, 1),
                        "scope": "vs_phase3_ready",
                    }
                )
                r["import_category"] = "DUPLICATE"
                phase4_new.remove(r)
                break

    final_ready = retained_p3 + phase4_new
    # Ensure unique ids
    final_by_id = {}
    for r in final_ready:
        final_by_id[r["id"]] = r
    final_ready = list(final_by_id.values())

    validation = validate_final(final_ready)

    unresolved = [
        r
        for r in recovered
        if r.get("import_category")
        in {"NEEDS_COORDINATES", "NEEDS_REVIEW", "COMING_SOON", "CLOSED"}
    ]

    final_n = len(final_ready)
    projected = 9056 + final_n
    headroom = 10000 - projected

    # Brand coverage after recovery
    live_brands = Counter(c.get("brand") for c in production_it)
    new_by_brand = Counter(r.get("brand") for r in final_ready)
    fx_meta = (discovery_meta.get("fitexpress") or {})
    icon_meta = (discovery_meta.get("icon") or {})
    orange_meta = (discovery_meta.get("orange") or {})

    fx_live = live_brands.get("Fit Express", 0)
    fx_new = new_by_brand.get("Fit Express", 0)
    fx_official = fx_meta.get("official_wp_count") or 69
    fx_cov = round(100.0 * (fx_live + fx_new) / fx_official, 1)

    icon_live = live_brands.get("Icon Palestre", 0)
    icon_new = new_by_brand.get("Icon Palestre", 0)
    icon_official = icon_meta.get("sitemap_clubish") or 47
    icon_cov = round(100.0 * (icon_live + icon_new) / icon_official, 1)

    # Stopping rule (user): Phase 5 ONLY if
    # - 100+ credible conventional chain gyms still recoverable, OR
    # - major national chain still materially absent (<50% of official estate), OR
    # - technical blocker on clearly large official estate (e.g. blocked API + ≥20 missing)
    still_unresolved_gap = [
        r
        for r in unresolved
        if r.get("brand") in {"Fit Express", "Icon Palestre", "Fitness Park", "FITINN"}
        and r.get("import_category") in {"NEEDS_COORDINATES", "NEEDS_REVIEW"}
    ]
    credible_left = len(still_unresolved_gap)
    # Prefer WP/staging estate sizes over inflated sitemap page counts
    icon_estate_for_gap = max(
        47,  # Phase 3 audited estate
        icon_live + len([r for r in recovered if r.get("brand") == "Icon Palestre"]),
    )
    # Recompute Icon coverage vs audited estate (not raw sitemap page count)
    icon_cov_audited = round(100.0 * (icon_live + icon_new) / max(icon_estate_for_gap, 1), 1)
    fx_missing = fx_official - fx_live - fx_new
    icon_missing = icon_estate_for_gap - icon_live - icon_new
    major_chain_absent = (fx_cov < 50) or (icon_cov_audited < 50)
    large_estate_blocker = (
        (fx_missing >= 20 and fx_cov < 75)
        or (icon_missing >= 20 and icon_cov_audited < 75)
        or (
            (orange_meta.get("getfit_site_status") or "").startswith("http_403")
            and (orange_meta.get("getfit_phase3_ready") or 0) == 0
            and (orange_meta.get("press_target_after_acquisition") or 0)
            - (orange_meta.get("production_live") or 0)
            >= 20
        )
    )
    phase5 = credible_left >= 100 or major_chain_absent or large_estate_blocker

    decision = (
        "ITALY PHASE 5 REQUIRED — MATERIAL MARKET GAP REMAINS"
        if phase5
        else "READY FOR ITALY COMPLETENESS MERGE"
    )

    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "production_catalog": 9056,
        "production_italy": 550,
        "phase3_ready_input": 22,
        "phase3_ready_retained": len(retained_p3),
        "phase3_demotions": demotions,
        "phase4_discovered": len(discovery),
        "phase4_newly_recovered_ready": len(phase4_new),
        "FINAL_NEW_READY": final_n,
        "ready_by_brand": dict(Counter(r.get("brand") for r in final_ready)),
        "phase4_new_by_brand": dict(Counter(r.get("brand") for r in phase4_new)),
        "projected_catalog": projected,
        "headroom_to_10k": headroom,
        "exceed_10k_at_new_ready": 945,
        "fitexpress_coverage": {
            "live": fx_live,
            "phase4_new_in_final": fx_new,
            "official_wp": fx_official,
            "coverage_pct_after": fx_cov,
            **fx_meta,
        },
        "icon_coverage": {
            "live": icon_live,
            "phase4_new_in_final": icon_new,
            "official_sitemap_clubish": icon_official,
            "audited_estate_estimate": icon_estate_for_gap,
            "coverage_pct_after_sitemap": icon_cov,
            "coverage_pct_after_audited": icon_cov_audited,
            "missing_vs_audited": icon_missing,
            **icon_meta,
        },
        "fitexpress_gap": {"missing_vs_wp": fx_missing, "coverage_pct_after": fx_cov},
        "orange_getfit": orange_meta,
        "unresolved_after_phase4": len(unresolved),
        "credible_conventional_still_recoverable_estimate": credible_left,
        "stopping_rule": {
            "credible_left": credible_left,
            "major_chain_absent": major_chain_absent,
            "large_estate_blocker": large_estate_blocker,
            "phase5_required": phase5,
        },
        "validation": {
            "dup_ids": len(validation["dup_ids"]),
            "already_in_production": len(validation["already_in_production"]),
            "gate_failures": len(validation["gate_failures"]),
            "ok": validation["ok"],
        },
        "phase5_required": phase5,
        "final_decision": decision,
    }

    # Write outputs
    (OUT / "ITALY_PHASE4_NEW_RECOVERED.json").write_text(
        json.dumps(phase4_new, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUT / "ITALY_PHASE4_GEOCODE_REVIEW.json").write_text(
        json.dumps(review, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUT / "ITALY_PHASE4_DUPLICATE_ANALYSIS.json").write_text(
        json.dumps(
            {
                **{k: (len(v) if isinstance(v, list) else v) for k, v in dup_report.items()},
                "details": dup_report,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    (OUT / "ITALY_COMPLETENESS_READY_TO_IMPORT.json").write_text(
        json.dumps(final_ready, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUT / "ITALY_PHASE4_COMPLETENESS_RECOVERY_REPORT.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    # Markdown report
    md = f"""# Italy Phase 4 — Final Completeness Recovery

Generated: {report['generated_at']}

**`src/data/centers.json` was NOT modified.** Production remains **9,056 / Italy 550**.

## Verdict

**{decision}**

## Catalog math

- Production catalog: **9056**
- Production Italy: **550** (unchanged)
- Phase 3 READY retained: **{len(retained_p3)}** (input 22; demotions {len(demotions)})
- Phase 4 newly recovered READY: **{len(phase4_new)}**
- FINAL_NEW_READY: **{final_n}**
- Projected catalog = 9056 + {final_n} = **{projected}**
- Headroom to 10k: **{headroom}** (exceeds at +945 FINAL_NEW_READY)

## Phase 3 READY retention

| Metric | Count |
|---|---:|
| Input READY | 22 |
| Retained | {len(retained_p3)} |
| Demoted | {len(demotions)} |

"""
    if demotions:
        md += "### Demotions\n\n"
        for d in demotions:
            md += f"- `{d['id']}` {d['name']}: {d['reason']}\n"
        md += "\n"
    else:
        md += "No demotions — all 22 Phase 3 READY rows retained.\n\n"

    md += f"""## Critical priorities

### Fit Express
- Official WP estate: **{fx_official}**
- Production live: **{fx_live}**
- Phase 4 new in final set: **{fx_new}**
- Coverage after Phase 4: **{fx_cov}%** of WP estate
- Note: {fx_meta.get('note', '')}

### Icon Palestre
- Brand identity: **Icon Palestre** (`iconpalestre.it`) — no generic Icon pollution
- Audited estate estimate: **{icon_estate_for_gap}** (sitemap clubish pages: {icon_official})
- Production live: **{icon_live}**
- Phase 4 new in final set: **{icon_new}**
- Coverage after Phase 4 (audited): **{icon_cov_audited}%**

### Orange / GetFIT
- Orange official locator: **{orange_meta.get('official_locator_count')}** (live production **{orange_meta.get('production_live')}**)
- GetFIT site: **{orange_meta.get('getfit_site_status')}**
- GetFIT Phase 3 READY acquired (not yet on Orange locator): **{orange_meta.get('getfit_acquired_not_on_orange_locator')}**
- Founder-retained GetFIT: **{orange_meta.get('getfit_founder_retained')}**
- Combined Orange live + acquired GetFIT READY: **{orange_meta.get('combined_orange_plus_acquired_getfit')}** (press target ~33)
- Handling: keep acquired clubs as **GetFIT** until Orange republishes; do not invent Orange rows

## FINAL READY by brand

| Brand | Count |
|---|---:|
"""
    for b, n in sorted(report["ready_by_brand"].items(), key=lambda x: (-x[1], x[0])):
        md += f"| {b} | {n} |\n"
    md += f"| **TOTAL** | **{final_n}** |\n\n"

    md += f"""## Phase 4 NEW recovered by brand

| Brand | Count |
|---|---:|
"""
    for b, n in sorted(report["phase4_new_by_brand"].items(), key=lambda x: (-x[1], x[0])):
        md += f"| {b} | {n} |\n"
    md += f"| **TOTAL** | **{len(phase4_new)}** |\n\n"

    md += f"""## Validation

- Dup IDs: **{validation['dup_ids'] and len(validation['dup_ids']) or 0}**
- Already in production: **{len(validation['already_in_production'])}**
- Gate failures: **{len(validation['gate_failures'])}**
- CLOSED/COMING_SOON in READY: **{len(validation['closed_or_coming'])}**
- Validation OK: **{validation['ok']}**

## Stopping rule

- Credible conventional still recoverable (unresolved FX/Icon/FP/FITINN leftovers): **{credible_left}**
- Major national chain materially absent (<50% coverage): **{major_chain_absent}**
- Technical blocker on clearly large official estate (≥20 missing & <75%): **{large_estate_blocker}**
- Threshold for Phase 5: 100+ recoverable OR major chain absent OR large-estate technical blocker
- Phase 5 required: **{phase5}**

## Files

- `scripts/italy-phase4-discover.py`
- `scripts/italy-phase4-consolidate.py`
- `data/italy/ITALY_PHASE4_COMPLETENESS_RECOVERY_REPORT.md`
- `data/italy/ITALY_PHASE4_COMPLETENESS_RECOVERY_REPORT.json`
- `data/italy/ITALY_PHASE4_NEW_RECOVERED.json`
- `data/italy/ITALY_PHASE4_GEOCODE_REVIEW.json`
- `data/italy/ITALY_PHASE4_DUPLICATE_ANALYSIS.json`
- `data/italy/ITALY_COMPLETENESS_READY_TO_IMPORT.json`
- `data/italy/Gymly_Italy_Final_Completeness.xlsx`

## FINAL: {decision}
"""
    (OUT / "ITALY_PHASE4_COMPLETENESS_RECOVERY_REPORT.md").write_text(md, encoding="utf-8")

    write_xlsx(retained_p3, phase4_new, final_ready, unresolved, report)

    log("FINAL_NEW_READY", final_n)
    log("projected", projected)
    log("validation_ok", validation["ok"])
    log("DECISION:", decision)
    if validation["gate_failures"]:
        log("gate failures sample", validation["gate_failures"][:5])


if __name__ == "__main__":
    main()
